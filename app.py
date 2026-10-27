import io
import json
import math
import os
import re
import secrets
import time
import zipfile
from datetime import date, datetime, timedelta
from functools import wraps

from flask import (Flask, abort, flash, g, redirect, render_template, request, send_file, session, url_for)
from openpyxl import load_workbook
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import check_password_hash, generate_password_hash

import excel
from bn import MONTHS_BN, bn_date, bn_num, bn_pct, bn_taka, month_label  # noqa: F401
from db import close_db, get_db, get_settings, init_db, save_setting
import payscale
from payroll import (ALLOWANCE_FIELDS, DEDUCTION_FIELDS, FIELD_LABELS, age_on, compute_components, finalize_totals,
                     FESTIVAL_MULTIPLIER, RELIGIONS, RUN_KINDS, arrears_already_given, arrears_final, generate_run,
                     LOAN_KINDS, SHORT_LABELS, group_by_department, insert_slip, loan_progress, record_loan_payments,
                     regenerate_slip, slip_for)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ভুল পাসওয়ার্ডের সীমা: একই IP + আইডি থেকে ১৫ মিনিটে ৫ বার
FAILED_LOGINS = {}
MAX_FAILED_LOGINS = 5
LOCKOUT_SECONDS = 15 * 60


def _secret_key(instance_dir):
    if os.environ.get("SECRET_KEY"):
        return os.environ["SECRET_KEY"]
    path = os.path.join(instance_dir, "secret_key")
    if not os.path.exists(path):
        with open(path, "w") as f:
            f.write(secrets.token_hex(32))
    with open(path) as f:
        return f.read().strip()


def create_app(database=None):
    instance_dir = os.path.join(BASE_DIR, "instance")
    os.makedirs(instance_dir, exist_ok=True)
    app = Flask(__name__)
    app.config.update(
        DATABASE=database or os.environ.get("DATABASE", os.path.join(instance_dir, "payroll.db")),
        SECRET_KEY=_secret_key(instance_dir),
        MAX_CONTENT_LENGTH=10 * 1024 * 1024,
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        # HTTPS এ চালালে SECURE_COOKIES=1 দিন
        SESSION_COOKIE_SECURE=os.environ.get("SECURE_COOKIES") == "1",
        PERMANENT_SESSION_LIFETIME=timedelta(hours=8),
    )
    if os.environ.get("BEHIND_PROXY") == "1":  # nginx / হোস্টিং প্রক্সির পেছনে
        app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
    app.teardown_appcontext(close_db)
    app.jinja_env.filters.update(bn=bn_num, taka=bn_taka, bndate=bn_date, pct=bn_pct, plain=plain_num)
    app.jinja_env.globals.update(MONTHS_BN=MONTHS_BN, month_label=month_label, FIELD_LABELS=FIELD_LABELS,
                                 loads=json.loads, RUN_KINDS=RUN_KINDS, RELIGIONS=RELIGIONS,
                                 FESTIVAL_MULTIPLIER=FESTIVAL_MULTIPLIER)
    with app.app_context():
        init_db()
    register_routes(app)
    return app


# ---------------------------------------------------------------- সহায়ক ফাংশন

BN_TO_EN = str.maketrans("০১২৩৪৫৬৭৮৯", "0123456789")


def to_int(value, default=0):
    if value is None:
        return default
    s = str(value).translate(BN_TO_EN).replace(",", "").strip()
    if s == "":
        return default
    try:
        return int(round(float(s)))
    except ValueError:
        return default


def plain_num(v):
    """ফর্মের ঘরে দেখানোর জন্য: 63960.0 → 63960, ফাঁকা → ''।"""
    if v is None or v == "":
        return ""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return v
    return int(f) if f.is_integer() else f


def to_float(value, default=0.0):
    s = str(value or "").translate(BN_TO_EN).replace(",", "").strip()
    try:
        return float(s) if s else default
    except ValueError:
        return default


def optional_num(value):
    """খালি ঘর = None (স্বয়ংক্রিয় হিসাব)।"""
    s = str(value if value is not None else "").translate(BN_TO_EN).replace(",", "").replace("/-", "").strip()
    return round(to_float(s), 2) if s else None


def clean_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    s = str(value or "").translate(BN_TO_EN).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s[:10], fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None


def current_user():
    uid = session.get("uid")
    if uid is None:
        return None
    if "user" not in g:
        g.user = get_db().execute("SELECT * FROM users WHERE id = ?", (uid,)).fetchone()
    return g.user


def login_required(view):
    @wraps(view)
    def wrapped(*a, **kw):
        if current_user() is None:
            return redirect(url_for("login", next=request.path))
        if current_user()["must_change"] and request.endpoint not in ("change_password", "logout"):
            flash("নিরাপত্তার জন্য প্রথমবার লগইনের পর পাসওয়ার্ড পরিবর্তন করুন।", "warn")
            return redirect(url_for("change_password"))
        return view(*a, **kw)
    return wrapped


def admin_required(view):
    @wraps(view)
    @login_required
    def wrapped(*a, **kw):
        if current_user()["role"] != "admin":
            abort(403)
        return view(*a, **kw)
    return wrapped


def xlsx_response(buf, filename):
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


EMPLOYEE_TEXT_FIELDS = ["code", "name_bn", "name_en", "father_name", "mother_name", "gender", "nid", "mobile",
                        "email", "address", "designation", "department", "employment_type", "status", "tin", "religion",
                        "bank_name", "bank_branch", "account_no"]
EMPLOYEE_DATE_FIELDS = ["dob", "join_date", "prl_date"]
OVERRIDE_FIELDS = ["ovr_house_rent", "ovr_medical", "ovr_education", "ovr_tiffin", "ovr_transport", "ovr_mobile"]
FLAG_FIELDS = ["govt_quarter", "inc_2026", "inc_2027", "pf", "tiffin", "transport", "wash"]
SPECIAL_RATES = ["auto", "0.10", "0.15", "0"]

# কর্মচারী নিজে যেসব তথ্য বদলাতে পারবেন (বেতন সংক্রান্ত কিছু নয়; পদবি দিয়ে বেতন হিসাব হয় না, গ্রেড দিয়ে হয়)
SELF_EDIT_FIELDS = [("name_bn", "নাম (বাংলায়)"), ("name_en", "নাম (ইংরেজিতে)"), ("designation", "পদবি"),
                    ("father_name", "পিতার নাম"),
                    ("mother_name", "মাতার নাম"), ("nid", "জাতীয় পরিচয়পত্র নং"), ("mobile", "মোবাইল"),
                    ("email", "ইমেইল"), ("address", "ঠিকানা"), ("bank_name", "ব্যাংকের নাম"),
                    ("bank_branch", "ব্যাংক শাখা"), ("account_no", "হিসাব নম্বর"), ("tin", "টিআইএন")]


# ইমপোর্টে নতুন কর্মচারীর ঘর খালি থাকলে
IMPORT_DEFAULTS = {"inc_2026": 1, "inc_2027": 1, "tiffin": 1, "transport": 1, "pf": 0, "wash": 0, "govt_quarter": 0,
                   "special_rate": "auto", "status": "active", "employment_type": "স্থায়ী", "children": 0}


RELIGION_ALIASES = {"ইসলাম": "ইসলাম", "মুসলিম": "ইসলাম", "মুসলমান": "ইসলাম", "islam": "ইসলাম", "muslim": "ইসলাম",
                    "হিন্দু": "হিন্দু", "সনাতন": "হিন্দু", "hindu": "হিন্দু", "বৌদ্ধ": "বৌদ্ধ", "buddhist": "বৌদ্ধ",
                    "buddhism": "বৌদ্ধ", "খ্রিস্টান": "খ্রিস্টান", "খ্রিষ্টান": "খ্রিস্টান", "christian": "খ্রিস্টান",
                    "অন্যান্য": "অন্যান্য", "other": "অন্যান্য"}


def normalize_religion(v):
    s = str(v or "").strip()
    return RELIGION_ALIASES.get(s.lower(), RELIGION_ALIASES.get(s)) if s else None


def parse_flag(v):
    s = str(v if v is not None else "").translate(BN_TO_EN).strip().lower()
    return s not in ("", "0", "না", "no", "n", "false", "none")


def employee_from_form(form):
    data = {f: (form.get(f) or "").strip() or None for f in EMPLOYEE_TEXT_FIELDS}
    data.update({f: clean_date(form.get(f)) for f in EMPLOYEE_DATE_FIELDS})
    data["status"] = data["status"] or "active"
    data["religion"] = normalize_religion(data["religion"])
    data["grade"] = to_int(form.get("grade"), None)
    data["basic"] = optional_num(form.get("basic")) or 0
    data["children"] = min(to_int(form.get("children")), 2)
    data.update({f: 1 if form.get(f) else 0 for f in FLAG_FIELDS})
    rate = form.get("special_rate") or "auto"
    data["special_rate"] = rate if rate in SPECIAL_RATES else "auto"
    data["new_basic_override"] = optional_num(form.get("new_basic_override"))
    data.update({f: optional_num(form.get(f)) for f in OVERRIDE_FIELDS})
    return data


SEPARATION_REASONS = ["অবসর", "অবসর-উত্তর ছুটি (PRL)", "পদত্যাগ", "বদলি", "চাকরিচ্যুত", "মৃত্যু", "অন্যান্য"]


def next_code():
    """পরের কর্মচারী আইডি: সবচেয়ে বড় NCC-নম্বর + ১।"""
    nums = [int(m.group(1)) for (code,) in get_db().execute("SELECT code FROM employees")
            if (m := re.fullmatch(r"NCC-(\d+)", code or ""))]
    return f"NCC-{(max(nums) + 1) if nums else 1:04d}"


def all_departments():
    return [r[0] for r in get_db().execute("SELECT name FROM departments ORDER BY id")]


def save_employee(data, emp_id=None):
    conn = get_db()
    if data.get("department"):
        conn.execute("INSERT OR IGNORE INTO departments (name) VALUES (?)", (data["department"],))
    cols = list(data.keys())
    if emp_id:
        conn.execute(f"UPDATE employees SET {', '.join(c + ' = ?' for c in cols)} WHERE id = ?",
                     [data[c] for c in cols] + [emp_id])
        conn.execute("UPDATE users SET username = ? WHERE employee_id = ?", (data["code"], emp_id))
    else:
        cur = conn.execute(f"INSERT INTO employees ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                           [data[c] for c in cols])
        emp_id = cur.lastrowid
        # প্রাথমিক পাসওয়ার্ড = কর্মচারী আইডি; প্রথম লগইনে পরিবর্তন বাধ্যতামূলক
        conn.execute("INSERT INTO users (username, password_hash, role, employee_id, must_change) "
                     "VALUES (?, ?, 'employee', ?, 1)", (data["code"], generate_password_hash(data["code"]), emp_id))
    return emp_id


def basic_range_warning(data):
    """৩০-০৬-২০২৬ এর মূল বেতন ২০১৫ স্কেলের ধাপের সাথে না মিললে সতর্কতা।"""
    if payscale.is_new_joiner(data):
        return
    steps = payscale.OLD.get(data["grade"])
    if steps and data["basic"] not in steps:
        flash(f"সতর্কতা: {bn_taka(data['basic'])} টাকা গ্রেড-{bn_num(data['grade'])} এর ২০১৫ স্কেলের কোনো ধাপের সাথে "
              f"মেলেনি ({bn_taka(steps[0])}–{bn_taka(steps[-1])})। হিসাব হবে, তবে সংখ্যাটি যাচাই করুন।", "warn")


def run_totals(kind=None):
    """প্রতিটি শিটের মোট; kind='salary' দিলে শুধু মাসিক বেতন শিট।"""
    return get_db().execute(
        """SELECT r.*, COUNT(p.id) AS headcount, COALESCE(SUM(p.gross),0) AS gross,
                  COALESCE(SUM(p.total_deduction),0) AS deduction, COALESCE(SUM(p.net),0) AS net
           FROM payroll_runs r LEFT JOIN payslips p ON p.run_id = r.id
           WHERE (? IS NULL OR r.kind = ?) AND r.manual = 0
           GROUP BY r.id ORDER BY r.year, r.month, r.id""", (kind, kind)).fetchall()


def bank_groups(run, slips):
    """ব্যাংক অনুযায়ী ভাগ: [(ব্যাংক, বিবরণী)], শেষে (None, হিসাব নম্বর নেই/অজানা)। খসড়া শিটে কর্মচারীর বর্তমান
    হিসাব নম্বর, চূড়ান্ত শিটে শিটের সময়ের নম্বর (খালি থাকলে বর্তমানটি)।"""
    rules = excel.parse_bank_rules(get_settings().get("bank_rules"))
    emps = {e["id"]: e for e in get_db().execute("SELECT id, account_no, bank_name FROM employees")}
    groups = {name: [] for name, _ in rules}
    groups[None] = []
    for s in slips:
        e = emps.get(s["employee_id"]) or {}
        acc = (e["account_no"] if e and (run["status"] == "draft" or not s["account_no"]) else s["account_no"]) or ""
        row = dict(s, account_no=acc)
        groups[excel.bank_of(acc, e["bank_name"] if e else None, rules)].append(row)
    return [(k, v) for k, v in groups.items() if v or k is not None]


def get_run_or_404(run_id):
    run = get_db().execute("SELECT * FROM payroll_runs WHERE id = ?", (run_id,)).fetchone()
    if run is None:
        abort(404)
    return run


def require_draft(run):
    if run["status"] != "draft":
        flash("চূড়ান্ত বেতন শিট পরিবর্তন করা যায় না। প্রয়োজনে আগে 'খসড়ায় ফেরত' দিন।", "error")
        return False
    return True


def pay_context(emp):
    """কর্মচারীর বেতন নির্ধারণ বিবরণী: পে ফিক্সেশন, পাঁচ তারিখের তুলনা, বকেয়া ও এ মাসের বেতন।"""
    today = date.today()
    pl = payscale.plan(emp)
    current = compute_components(emp, get_settings(), today.year, today.month) if "error" not in pl else None
    return {"plan": pl, "compare": payscale.comparison(emp) if "res" in pl else [],
            "arrears": payscale.arrears(emp), "current": current, "today": today, "phase": payscale.phase_of(today.year, today.month),
            "age": age_on(emp["dob"], today), "prl": payscale.prl_date(emp["dob"]), "special_rate": payscale.special_rate(emp)
            if emp["grade"] else 0}


def employee_loans(emp_id):
    loans = get_db().execute("SELECT * FROM loans WHERE employee_id = ? ORDER BY status, id", (emp_id,)).fetchall()
    return [{"loan": l, **loan_progress(l)} for l in loans]


def loan_from_form(form):
    ym = (form.get("start_ym") or "").strip()
    return {"kind": (form.get("kind") or "অন্যান্য ঋণ").strip(), "installment": round(to_float(form.get("installment")), 2),
            "total_installments": to_int(form.get("total_installments"), None) or None,
            "paid_before": to_int(form.get("paid_before")), "principal": optional_num(form.get("principal")),
            "start_ym": ym if re.fullmatch(r"\d{4}-\d{2}", ym) else None, "note": (form.get("note") or "").strip() or None}


def items_text(emp_id, kind):
    rows = get_db().execute("SELECT name, amount FROM employee_items WHERE employee_id = ? AND kind = ? ORDER BY id",
                            (emp_id, kind)).fetchall()
    return "; ".join(f"{r['name']}={plain_num(r['amount'])}" for r in rows) or None


def parse_items(text):
    """'মোবাইল ভাতা=1500; আয়কর=500' → [(নাম, টাকা)]"""
    out = []
    for part in re.split(r"[;\n]", str(text or "")):
        if "=" in part:
            name, amt = part.rsplit("=", 1)
            if name.strip():
                out.append((name.strip(), round(to_float(amt), 2)))
    return out


def employee_export_row(e):
    d = dict(e)
    d["allow_items"], d["ded_items"] = items_text(e["id"], "allowance"), items_text(e["id"], "deduction")
    for i, kind in enumerate(excel.IMPORT_LOAN_KINDS):
        loan = get_db().execute("SELECT * FROM loans WHERE employee_id = ? AND kind = ? AND status = 'active' "
                                "ORDER BY id DESC", (e["id"], kind)).fetchone()
        if loan:
            pr = loan_progress(loan)
            d[f"loan{i}_inst"], d[f"loan{i}_total"], d[f"loan{i}_paid"] = loan["installment"], loan["total_installments"], pr["paid"]
    return d


def import_extras(emp_id, raw):
    """ইমপোর্টের অতিরিক্ত কলাম: ভাতা/কর্তনের তালিকা ও ঋণ (কলাম খালি থাকলে আগের তথ্য অপরিবর্তিত)।"""
    conn = get_db()
    for col, kind in [("allow_items", "allowance"), ("ded_items", "deduction")]:
        if raw.get(col) not in (None, ""):
            conn.execute("DELETE FROM employee_items WHERE employee_id = ? AND kind = ?", (emp_id, kind))
            for name, amt in parse_items(raw[col]):
                conn.execute("INSERT INTO employee_items (employee_id, kind, name, amount) VALUES (?, ?, ?, ?)",
                             (emp_id, kind, name, amt))
    for i, kind in enumerate(excel.IMPORT_LOAN_KINDS):
        inst = raw.get(f"loan{i}_inst")
        if inst in (None, ""):
            continue
        inst = round(to_float(inst), 2)
        total = to_int(raw.get(f"loan{i}_total"), None) or None
        paid = raw.get(f"loan{i}_paid")
        loan = conn.execute("SELECT * FROM loans WHERE employee_id = ? AND kind = ? AND status = 'active' ORDER BY id DESC",
                            (emp_id, kind)).fetchone()
        if inst <= 0:
            if loan:
                conn.execute("UPDATE loans SET status = 'closed' WHERE id = ?", (loan["id"],))
        elif loan:
            # পরিশোধিত কিস্তির ঘরে দেওয়া সংখ্যা = আগে পরিশোধিত + এই সফটওয়্যারে চূড়ান্ত শিটে কাটা
            final = loan_progress(loan)["paid"] - loan["paid_before"]
            paid_before = max(to_int(paid) - final, 0) if paid not in (None, "") else loan["paid_before"]
            conn.execute("UPDATE loans SET installment = ?, total_installments = ?, paid_before = ? WHERE id = ?",
                         (inst, total, paid_before, loan["id"]))
        else:
            conn.execute("INSERT INTO loans (employee_id, kind, installment, total_installments, paid_before) "
                         "VALUES (?, ?, ?, ?, ?)", (emp_id, kind, inst, total, to_int(paid)))


def employee_slips(emp_id, final_only):
    sql = ("SELECT p.*, r.year, r.month, r.status, r.kind, r.title FROM payslips p JOIN payroll_runs r ON r.id = p.run_id "
           "WHERE p.employee_id = ?" + (" AND r.status = 'final'" if final_only else "") +
           " ORDER BY r.year DESC, r.month DESC, r.id DESC")
    return get_db().execute(sql, (emp_id,)).fetchall()


# ---------------------------------------------------------------- রাউট

def register_routes(app):

    @app.before_request
    def csrf_protect():
        if "csrf" not in session:
            session["csrf"] = secrets.token_hex(16)
        if request.method == "POST" and request.form.get("csrf") != session["csrf"]:
            abort(400, "অবৈধ অনুরোধ (CSRF)। পাতাটি রিফ্রেশ করে আবার চেষ্টা করুন।")

    @app.context_processor
    def inject():
        return {"user": current_user(), "csrf": session.get("csrf", ""), "settings": get_settings()}

    # ---------- লগইন

    @app.route("/")
    def index():
        user = current_user()
        if user is None:
            return redirect(url_for("login"))
        return redirect(url_for("admin_dashboard" if user["role"] == "admin" else "my_dashboard"))

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if request.method == "POST":
            username = request.form.get("username", "").strip()
            key = (request.remote_addr, username.lower())
            now = time.time()
            fails = [t for t in FAILED_LOGINS.get(key, []) if now - t < LOCKOUT_SECONDS]
            if len(fails) >= MAX_FAILED_LOGINS:
                flash("বারবার ভুল পাসওয়ার্ডের কারণে ১৫ মিনিটের জন্য লগইন বন্ধ। পরে চেষ্টা করুন।", "error")
                return render_template("login.html"), 429
            user = get_db().execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
            if user and check_password_hash(user["password_hash"], request.form.get("password", "")):
                FAILED_LOGINS.pop(key, None)
                if user["employee_id"]:
                    emp = get_db().execute("SELECT status FROM employees WHERE id = ?", (user["employee_id"],)).fetchone()
                    if emp and emp["status"] != "active":
                        flash("এই অ্যাকাউন্টটি নিষ্ক্রিয়। প্রশাসনের সাথে যোগাযোগ করুন।", "error")
                        return render_template("login.html")
                session.clear()
                session.permanent = True
                session["uid"] = user["id"]
                nxt = request.args.get("next", "")
                return redirect(nxt if nxt.startswith("/") and not nxt.startswith("//") else url_for("index"))
            FAILED_LOGINS[key] = fails + [now]
            flash("আইডি অথবা পাসওয়ার্ড সঠিক নয়।", "error")
        return render_template("login.html")

    @app.route("/logout", methods=["POST"])
    def logout():
        session.clear()
        return redirect(url_for("login"))

    @app.route("/password", methods=["GET", "POST"])
    @login_required
    def change_password():
        if request.method == "POST":
            user = current_user()
            new = request.form.get("new", "")
            if not check_password_hash(user["password_hash"], request.form.get("old", "")):
                flash("বর্তমান পাসওয়ার্ড সঠিক নয়।", "error")
            elif len(new) < 6:
                flash("নতুন পাসওয়ার্ড কমপক্ষে ৬ অক্ষরের হতে হবে।", "error")
            elif new != request.form.get("confirm"):
                flash("নতুন পাসওয়ার্ড দুইবার একই হয়নি।", "error")
            elif new == user["username"]:
                flash("আইডি আর পাসওয়ার্ড একই রাখা যাবে না।", "error")
            else:
                get_db().execute("UPDATE users SET password_hash = ?, must_change = 0 WHERE id = ?",
                                 (generate_password_hash(new), user["id"]))
                get_db().commit()
                flash("পাসওয়ার্ড পরিবর্তন হয়েছে।", "ok")
                return redirect(url_for("index"))
        return render_template("password.html")

    # ---------- অ্যাডমিন ড্যাশবোর্ড

    @app.route("/admin")
    @admin_required
    def admin_dashboard():
        conn = get_db()
        runs = run_totals("salary")
        stats = conn.execute(
            "SELECT COUNT(*) AS total, SUM(status = 'active') AS active, "
            "COALESCE(SUM(CASE WHEN status = 'active' THEN basic END), 0) AS basic_sum FROM employees").fetchone()
        depts = conn.execute(
            "SELECT COALESCE(department, '—') AS department, COUNT(*) AS n FROM employees WHERE status = 'active' "
            "GROUP BY department ORDER BY n DESC").fetchall()
        today = date.today()
        this_run = conn.execute("SELECT * FROM payroll_runs WHERE year = ? AND month = ? AND kind = 'salary' AND manual = 0",
                                (today.year, today.month)).fetchone()
        changes = conn.execute(
            "SELECT c.*, e.name_bn, e.code FROM employee_changes c JOIN employees e ON e.id = c.employee_id "
            "ORDER BY c.id DESC LIMIT 10").fetchall()
        return render_template("admin/dashboard.html", stats=stats, runs=runs[-12:], last=runs[-1] if runs else None,
                               prev=runs[-2] if len(runs) > 1 else None, depts=depts, this_run=this_run, today=today,
                               changes=changes, phase=payscale.phase_of(today.year, today.month))

    # ---------- কর্মচারী

    @app.route("/admin/employees")
    @admin_required
    def employees():
        q = request.args.get("q", "").strip()
        dept = request.args.get("dept", "")
        status = request.args.get("status", "active")
        sql, args = "SELECT * FROM employees WHERE 1=1", []
        if q:
            sql += " AND (code LIKE ? OR name_bn LIKE ? OR name_en LIKE ? OR designation LIKE ? OR mobile LIKE ?)"
            args += [f"%{q}%"] * 5
        if dept:
            sql += " AND department = ?"
            args.append(dept)
        if status:
            sql += " AND status = ?"
            args.append(status)
        religion = request.args.get("religion", "")
        if religion == "none":
            sql += " AND (religion IS NULL OR religion = '')"
        elif religion:
            sql += " AND religion = ?"
            args.append(religion)
        rows = get_db().execute(sql + " ORDER BY code", args).fetchall()
        depts = all_departments()
        counts = dict(get_db().execute("SELECT status, COUNT(*) FROM employees GROUP BY status").fetchall())
        no_religion = get_db().execute("SELECT COUNT(*) FROM employees WHERE status = 'active' AND "
                                       "(religion IS NULL OR religion = '')").fetchone()[0]
        return render_template("admin/employees.html", rows=rows, q=q, dept=dept, status=status, depts=depts,
                               religion=religion, no_religion=no_religion, today_iso=date.today().isoformat(),
                               counts=counts, next_code=next_code())

    def employee_form_context(emp=None):
        conn = get_db()
        depts = all_departments()
        desigs = [r[0] for r in conn.execute(
            "SELECT DISTINCT designation FROM employees WHERE designation IS NOT NULL ORDER BY designation")]
        return {"emp": emp, "scales": payscale.OLD, "depts": depts, "desigs": desigs, "new_code": next_code()}

    @app.route("/admin/employees/new", methods=["GET", "POST"])
    @admin_required
    def employee_new():
        if request.method == "POST":
            data = employee_from_form(request.form)
            data["code"] = data["code"] or next_code()
            if not data["name_bn"]:
                flash("নাম অবশ্যই দিতে হবে।", "error")
            elif get_db().execute("SELECT 1 FROM users WHERE username = ?", (data["code"],)).fetchone():
                flash(f"আইডি {data['code']} আগে থেকেই আছে। পরের খালি আইডি: {next_code()}", "error")
            else:
                emp_id = save_employee(data)
                get_db().commit()
                basic_range_warning(data)
                flash(f"কর্মচারী যুক্ত হয়েছে। লগইন আইডি ও প্রাথমিক পাসওয়ার্ড: {data['code']}", "ok")
                return redirect(url_for("employee_view", emp_id=emp_id))
            return render_template("admin/employee_form.html", **employee_form_context(data))
        return render_template("admin/employee_form.html", **employee_form_context())

    @app.route("/admin/employees/<int:emp_id>/edit", methods=["GET", "POST"])
    @admin_required
    def employee_edit(emp_id):
        emp = get_db().execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone() or abort(404)
        if request.method == "POST":
            data = employee_from_form(request.form)
            clash = get_db().execute("SELECT 1 FROM users WHERE username = ? AND COALESCE(employee_id, 0) != ?",
                                     (data["code"], emp_id)).fetchone()
            if not data["code"] or not data["name_bn"]:
                flash("আইডি ও নাম অবশ্যই দিতে হবে।", "error")
            elif clash:
                flash("এই আইডি অন্য কারো নামে আছে।", "error")
            else:
                save_employee(data, emp_id)
                get_db().commit()
                basic_range_warning(data)
                flash("তথ্য হালনাগাদ হয়েছে।", "ok")
                return redirect(url_for("employee_view", emp_id=emp_id))
            data["id"] = emp_id
            return render_template("admin/employee_form.html", **employee_form_context(data))
        return render_template("admin/employee_form.html", **employee_form_context(emp))

    @app.route("/admin/employees/<int:emp_id>/manual", methods=["GET", "POST"])
    @admin_required
    def manual_slip(emp_id):
        """যে মাসের বেতন শিট সফটওয়্যারে নেই (সফটওয়্যার চালুর আগের মাস), সে মাসের বেতন হাতে এন্ট্রি।
        চূড়ান্ত হিসেবে সংরক্ষিত হয়, তাই কর্মচারী নিজেও বেতন বিবরণীতে দেখতে পান; মাসিক তালিকা/রিপোর্টে আসে না।"""
        conn = get_db()
        emp = conn.execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone() or abort(404)
        src = request.form if request.method == "POST" else request.args
        year, month = to_int(src.get("year")), to_int(src.get("month"))
        kind = src.get("kind") if src.get("kind") in RUN_KINDS else "salary"
        title = (src.get("title") or "").strip() or (None if kind == "salary" else RUN_KINDS[kind])
        valid = 2000 <= year <= 2100 and 1 <= month <= 12
        if request.method == "POST":
            if not valid:
                flash("সঠিক মাস ও বছর দিন।", "error")
                return redirect(url_for("manual_slip", emp_id=emp_id))
            official = conn.execute(
                "SELECT r.id FROM payroll_runs r JOIN payslips p ON p.run_id = r.id WHERE r.year = ? AND r.month = ? "
                "AND r.kind = ? AND r.manual = 0 AND p.employee_id = ?", (year, month, kind, emp_id)).fetchone()
            if official:
                flash(f"{month_label(year, month)} এর {RUN_KINDS[kind]} সফটওয়্যারের শিটেই আছে — সেখান থেকে সংশোধন করুন।",
                      "error")
                return redirect(url_for("payroll_view", run_id=official["id"]))
            s = {f: round(to_float(request.form.get(f)), 2) for f in ALLOWANCE_FIELDS + DEDUCTION_FIELDS}
            for part in ("other_allowance", "other_deduction"):
                names, amounts = request.form.getlist(f"{part}_name"), request.form.getlist(f"{part}_amount")
                s[part] = [{"name": n.strip(), "amount": round(to_float(a), 2)}
                           for n, a in zip(names, amounts) if n.strip() and to_float(a)]
            finalize_totals(s)
            if not s["gross"]:
                flash("কমপক্ষে একটি ভাতার টাকা দিন।", "error")
                return redirect(url_for("manual_slip", emp_id=emp_id, year=year, month=month, kind=kind))
            run = conn.execute("SELECT * FROM payroll_runs WHERE year = ? AND month = ? AND kind = ? AND manual = 1 "
                               "AND COALESCE(title, '') = ?", (year, month, kind, title or "")).fetchone()
            if run is None:
                cur = conn.execute(
                    "INSERT INTO payroll_runs (year, month, kind, title, status, manual, note, finalized_at) "
                    "VALUES (?, ?, ?, ?, 'final', 1, 'পূর্ববর্তী মাস — হাতে এন্ট্রি', CURRENT_TIMESTAMP)",
                    (year, month, kind, title))
                run_id = cur.lastrowid
            else:
                run_id = run["id"]
            # একই মাসের আগের এন্ট্রি (বা সংশোধনের সময় মাস বদলালে পুরনোটি) সরিয়ে নতুনটি
            conn.execute("DELETE FROM payslips WHERE run_id = ? AND employee_id = ?", (run_id, emp_id))
            old = conn.execute("SELECT p.id, p.run_id FROM payslips p JOIN payroll_runs r ON r.id = p.run_id "
                               "WHERE p.id = ? AND p.employee_id = ? AND r.manual = 1",
                               (to_int(request.form.get("slip_id")), emp_id)).fetchone()
            if old:
                conn.execute("DELETE FROM payslips WHERE id = ?", (old["id"],))
                conn.execute("DELETE FROM payroll_runs WHERE id = ? AND NOT EXISTS (SELECT 1 FROM payslips WHERE run_id = ?)",
                             (old["run_id"], old["run_id"]))
            s.update(phase="হাতে এন্ট্রি", ref_basic=None)
            insert_slip(run_id, emp, s)
            conn.commit()
            flash(f"{month_label(year, month)} — {title or RUN_KINDS[kind]}: {emp['name_bn']} এর বেতন সংরক্ষিত হয়েছে।", "ok")
            return redirect(url_for("manual_slip", emp_id=emp_id))

        # ফর্মে আগে থেকে বসানো মান: আগের এন্ট্রি সম্পাদনা, অথবা ক্যালকুলেটর থেকে হিসাব
        slip, note = None, None
        if src.get("slip"):
            slip = conn.execute("SELECT p.*, r.year, r.month, r.kind, r.title FROM payslips p JOIN payroll_runs r "
                                "ON r.id = p.run_id WHERE p.id = ? AND p.employee_id = ? AND r.manual = 1",
                                (to_int(src.get("slip")), emp_id)).fetchone()
            if slip:
                year, month, kind, title = slip["year"], slip["month"], slip["kind"], slip["title"]
        elif valid and kind == "salary" and src.get("calc"):
            pseudo = {"id": 0, "year": year, "month": month, "kind": "salary", "include_festival": 0,
                      "include_boishakhi": 0, "include_arrears": 0}
            try:
                calc = compute_components(emp, get_settings(), year, month, run=pseudo)
            except Exception:
                calc = None
            if calc:
                calc["other_deduction"] = [{"name": x["name"].split(" কিস্তি")[0], "amount": x["amount"]}
                                           for x in calc["other_deduction"]]
                slip = dict(calc, other_allowance=json.dumps(calc["other_allowance"], ensure_ascii=False),
                            other_deduction=json.dumps(calc["other_deduction"], ensure_ascii=False))
                note = "কর্মচারীর বর্তমান তথ্য ও বেতন স্কেল অনুযায়ী হিসাব বসানো হয়েছে — সে মাসের প্রকৃত বেতনের সাথে মিলিয়ে প্রয়োজনে সংশোধন করুন।"
            else:
                note = "কর্মচারীর গ্রেড/মূল বেতন অসম্পূর্ণ, তাই হিসাব বসানো যায়নি — টাকাগুলো হাতে লিখুন।"
        if slip is None:
            slip = {f: 0 for f in ALLOWANCE_FIELDS + DEDUCTION_FIELDS}
            slip.update(other_allowance="[]", other_deduction="[]")
        entries = conn.execute(
            "SELECT p.*, r.year, r.month, r.kind, r.title FROM payslips p JOIN payroll_runs r ON r.id = p.run_id "
            "WHERE p.employee_id = ? AND r.manual = 1 ORDER BY r.year DESC, r.month DESC, r.id DESC", (emp_id,)).fetchall()
        today = date.today()
        return render_template("admin/manual_slip.html", emp=emp, slip=slip, entries=entries, note=note,
                               year=year if valid else today.year, month=month if valid else today.month,
                               kind=kind, title=title or "", allowance_fields=ALLOWANCE_FIELDS,
                               deduction_fields=DEDUCTION_FIELDS, run_kinds=RUN_KINDS,
                               editing=slip["id"] if src.get("slip") and "id" in slip.keys() else None)

    @app.route("/admin/manual/<int:slip_id>/delete", methods=["POST"])
    @admin_required
    def manual_slip_delete(slip_id):
        conn = get_db()
        slip = conn.execute("SELECT p.* FROM payslips p JOIN payroll_runs r ON r.id = p.run_id "
                            "WHERE p.id = ? AND r.manual = 1", (slip_id,)).fetchone() or abort(404)
        conn.execute("DELETE FROM payslips WHERE id = ?", (slip_id,))
        conn.execute("DELETE FROM payroll_runs WHERE id = ? AND manual = 1 AND "
                     "NOT EXISTS (SELECT 1 FROM payslips WHERE run_id = ?)", (slip["run_id"], slip["run_id"]))
        conn.commit()
        flash("এন্ট্রি মুছে ফেলা হয়েছে।", "ok")
        return redirect(url_for("manual_slip", emp_id=slip["employee_id"]))

    @app.route("/admin/employees/<int:emp_id>")
    @admin_required
    def employee_view(emp_id):
        conn = get_db()
        emp = conn.execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone() or abort(404)
        items = conn.execute("SELECT * FROM employee_items WHERE employee_id = ? ORDER BY kind, id", (emp_id,)).fetchall()
        account = conn.execute("SELECT * FROM users WHERE employee_id = ?", (emp_id,)).fetchone()
        changes = conn.execute("SELECT * FROM employee_changes WHERE employee_id = ? ORDER BY id DESC LIMIT 20",
                               (emp_id,)).fetchall()
        return render_template("admin/employee_view.html", emp=emp, items=items, account=account, changes=changes,
                               loans=employee_loans(emp_id), LOAN_KINDS=LOAN_KINDS,
                               slips=employee_slips(emp_id, final_only=False), **pay_context(emp))

    @app.route("/admin/employees/<int:emp_id>/items", methods=["POST"])
    @admin_required
    def employee_item_add(emp_id):
        name = request.form.get("name", "").strip()
        kind = request.form.get("kind")
        if name and kind in ("allowance", "deduction"):
            get_db().execute("INSERT INTO employee_items (employee_id, kind, name, amount) VALUES (?, ?, ?, ?)",
                             (emp_id, kind, name, round(to_float(request.form.get("amount")), 2)))
            get_db().commit()
            flash("যুক্ত হয়েছে।", "ok")
        return redirect(url_for("employee_view", emp_id=emp_id) + "#salary")

    @app.route("/admin/items/<int:item_id>/delete", methods=["POST"])
    @admin_required
    def employee_item_delete(item_id):
        item = get_db().execute("SELECT * FROM employee_items WHERE id = ?", (item_id,)).fetchone() or abort(404)
        get_db().execute("DELETE FROM employee_items WHERE id = ?", (item_id,))
        get_db().commit()
        return redirect(url_for("employee_view", emp_id=item["employee_id"]) + "#salary")

    @app.route("/admin/employees/<int:emp_id>/reset-password", methods=["POST"])
    @admin_required
    def employee_reset_password(emp_id):
        emp = get_db().execute("SELECT code FROM employees WHERE id = ?", (emp_id,)).fetchone() or abort(404)
        get_db().execute("UPDATE users SET password_hash = ?, must_change = 1 WHERE employee_id = ?",
                         (generate_password_hash(emp["code"]), emp_id))
        get_db().commit()
        flash(f"পাসওয়ার্ড রিসেট হয়েছে। নতুন প্রাথমিক পাসওয়ার্ড: {emp['code']}", "ok")
        return redirect(url_for("employee_view", emp_id=emp_id))

    @app.route("/admin/employees/<int:emp_id>/deactivate", methods=["POST"])
    @admin_required
    def employee_deactivate(emp_id):
        """'মুছুন' = নিষ্ক্রিয় তালিকায় সরানো (অবসর/পদত্যাগ ইত্যাদি); বেতনের পুরনো রেকর্ড থেকে যায়।"""
        emp = get_db().execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone() or abort(404)
        reason = request.form.get("reason") or "অন্যান্য"
        when = clean_date(request.form.get("date")) or date.today().isoformat()
        get_db().execute("UPDATE employees SET status = 'inactive', separation_reason = ?, separation_date = ? WHERE id = ?",
                         (reason, when, emp_id))
        get_db().commit()
        flash(f"{emp['name_bn']} নিষ্ক্রিয় তালিকায় সরানো হয়েছে ({reason}, {bn_date(when)})। পরের বেতন শিটে আসবেন না, "
              "লগইনও করতে পারবেন না; পুরনো বেতনের রেকর্ড থাকবে।", "ok")
        return redirect(url_for("employees", status="inactive"))

    @app.route("/admin/employees/<int:emp_id>/activate", methods=["POST"])
    @admin_required
    def employee_activate(emp_id):
        emp = get_db().execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone() or abort(404)
        get_db().execute("UPDATE employees SET status = 'active', separation_reason = NULL, separation_date = NULL "
                         "WHERE id = ?", (emp_id,))
        get_db().commit()
        flash(f"{emp['name_bn']} আবার সক্রিয় করা হয়েছে।", "ok")
        return redirect(url_for("employee_view", emp_id=emp_id))

    @app.route("/admin/employees/<int:emp_id>/delete", methods=["POST"])
    @admin_required
    def employee_delete(emp_id):
        """স্থায়ীভাবে মোছা — শুধু ভুল করে যুক্ত কর্মচারীর জন্য, যার কোনো বেতন শিটে নাম নেই।"""
        conn = get_db()
        emp = conn.execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone() or abort(404)
        if conn.execute("SELECT 1 FROM payslips WHERE employee_id = ?", (emp_id,)).fetchone():
            flash("এই কর্মচারীর বেতনের রেকর্ড আছে, তাই স্থায়ীভাবে মোছা যাবে না। নিষ্ক্রিয় করুন।", "error")
            return redirect(url_for("employee_view", emp_id=emp_id))
        conn.execute("DELETE FROM users WHERE employee_id = ?", (emp_id,))
        conn.execute("DELETE FROM employees WHERE id = ?", (emp_id,))
        conn.commit()
        flash(f"{emp['name_bn']} ({emp['code']}) স্থায়ীভাবে মুছে ফেলা হয়েছে।", "ok")
        return redirect(url_for("employees"))

    @app.route("/admin/employees/<int:emp_id>/loans", methods=["POST"])
    @admin_required
    def loan_add(emp_id):
        get_db().execute("SELECT 1 FROM employees WHERE id = ?", (emp_id,)).fetchone() or abort(404)
        d = loan_from_form(request.form)
        if d["installment"] <= 0:
            flash("কিস্তির টাকা দিন।", "error")
        else:
            cols = list(d)
            get_db().execute(f"INSERT INTO loans (employee_id, {', '.join(cols)}) VALUES (?, {', '.join('?' * len(cols))})",
                             [emp_id] + [d[c] for c in cols])
            get_db().commit()
            flash(f"{d['kind']} যুক্ত হয়েছে — প্রতি মাসের বেতন শিটে {bn_taka(d['installment'])} টাকা করে কাটা হবে।", "ok")
        return redirect(url_for("employee_view", emp_id=emp_id) + "#loans")

    @app.route("/admin/loans/<int:loan_id>", methods=["POST"])
    @admin_required
    def loan_update(loan_id):
        loan = get_db().execute("SELECT * FROM loans WHERE id = ?", (loan_id,)).fetchone() or abort(404)
        action = request.form.get("action")
        if action == "delete":
            if get_db().execute("SELECT 1 FROM loan_payments p JOIN payroll_runs r ON r.id = p.run_id "
                                "WHERE p.loan_id = ? AND r.status = 'final'", (loan_id,)).fetchone():
                flash("এই ঋণের কিস্তি চূড়ান্ত বেতন শিটে কাটা হয়েছে, তাই মোছা যাবে না — 'বন্ধ করুন' ব্যবহার করুন।", "error")
            else:
                get_db().execute("DELETE FROM loans WHERE id = ?", (loan_id,))
                flash("ঋণ মুছে ফেলা হয়েছে।", "ok")
        elif action in ("close", "reopen"):
            get_db().execute("UPDATE loans SET status = ? WHERE id = ?", ("closed" if action == "close" else "active", loan_id))
            flash("ঋণ বন্ধ করা হয়েছে — আর কিস্তি কাটা হবে না।" if action == "close" else "ঋণ আবার চালু করা হয়েছে।", "ok")
        else:
            d = loan_from_form(request.form)
            get_db().execute("UPDATE loans SET kind = ?, installment = ?, total_installments = ?, paid_before = ?, "
                             "principal = ?, start_ym = ?, note = ? WHERE id = ?",
                             [d[c] for c in ["kind", "installment", "total_installments", "paid_before", "principal",
                                             "start_ym", "note"]] + [loan_id])
            flash("ঋণের তথ্য হালনাগাদ হয়েছে। খসড়া শিটে প্রভাব দেখতে সেই শিটে 'সবার পুনঃহিসাব' চাপুন।", "ok")
        get_db().commit()
        return redirect(url_for("employee_view", emp_id=loan["employee_id"]) + "#loans")

    @app.route("/admin/employees/<int:emp_id>/history.xlsx")
    @admin_required
    def employee_history_xlsx(emp_id):
        emp = get_db().execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone() or abort(404)
        slips = list(reversed(employee_slips(emp_id, final_only=False)))
        return xlsx_response(excel.employee_history_workbook(emp, slips, get_settings()), f"salary-{emp['code']}.xlsx")

    @app.route("/admin/employees/export.xlsx")
    @admin_required
    def employees_export():
        rows = [employee_export_row(e) for e in get_db().execute("SELECT * FROM employees ORDER BY code")]
        return xlsx_response(excel.employees_workbook(rows, get_settings()), "employees.xlsx")

    @app.route("/admin/employees/template.xlsx")
    @admin_required
    def employees_template():
        return xlsx_response(excel.employees_workbook([], get_settings(), template=True), "employee-import-template.xlsx")

    @app.route("/admin/employees/import", methods=["GET", "POST"])
    @admin_required
    def employees_import():
        if request.method == "POST":
            f = request.files.get("file")
            if not f or not f.filename.lower().endswith(".xlsx"):
                flash("একটি .xlsx ফাইল নির্বাচন করুন।", "error")
                return redirect(request.url)
            try:
                ws = load_workbook(f, data_only=True).active
            except Exception:
                flash("ফাইলটি পড়া যায়নি। টেমপ্লেট অনুযায়ী .xlsx ফাইল দিন।", "error")
                return redirect(request.url)
            header_to_field = {h.rstrip("*").strip(): fld for fld, h in excel.EMPLOYEE_COLUMNS}
            header = [header_to_field.get(str(c.value or "").rstrip("*").strip()) for c in ws[1]]
            if "name_bn" not in header:
                flash("প্রথম সারিতে টেমপ্লেটের শিরোনাম (নাম) পাওয়া যায়নি।", "error")
                return redirect(request.url)
            conn = get_db()
            added = updated = 0
            errors = []
            for r_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                raw = {fld: v for fld, v in zip(header, row) if fld}
                if not any(v not in (None, "") for v in raw.values()):
                    continue
                if not raw.get("name_bn"):
                    errors.append(f"সারি {bn_num(r_idx)}: নাম নেই")
                    continue
                code = str(raw.get("code") or "").strip() or next_code()
                raw["code"] = code
                existing = conn.execute("SELECT * FROM employees WHERE code = ?", (code,)).fetchone()
                form = dict(existing) if existing else dict(IMPORT_DEFAULTS)
                form.update({k: v for k, v in raw.items() if v not in (None, "")})
                for k in FLAG_FIELDS:
                    form[k] = "1" if parse_flag(form.get(k)) else ""
                form = {k: ("" if v is None else v if isinstance(v, (date, datetime)) else str(v)) for k, v in form.items()}
                data = employee_from_form(form)
                if existing:
                    emp_id = save_employee(data, existing["id"])
                    updated += 1
                elif conn.execute("SELECT 1 FROM users WHERE username = ?", (code,)).fetchone():
                    errors.append(f"সারি {bn_num(r_idx)}: আইডি '{code}' ব্যবহারকারী নাম হিসেবে আগেই আছে")
                    continue
                else:
                    emp_id = save_employee(data)
                    added += 1
                import_extras(emp_id, raw)
                if "error" in payscale.plan(data):
                    errors.append(f"সারি {bn_num(r_idx)} ({data['name_bn']}): বেতন হিসাব হয়নি — {payscale.plan(data)['error']}")
            conn.commit()
            flash(f"ইমপোর্ট সম্পন্ন: নতুন {bn_num(added)} জন, হালনাগাদ {bn_num(updated)} জন।", "ok")
            for e in errors[:30]:
                flash(e, "warn")
            return redirect(url_for("employees"))
        return render_template("admin/import.html")

    @app.route("/admin/employees/bulk-status", methods=["POST"])
    @admin_required
    def employees_bulk_status():
        """তালিকা থেকে টিক দেওয়া কর্মচারীদের একসাথে নিষ্ক্রিয় (অবসর/পদত্যাগ/বাতিল ...) বা আবার সক্রিয় করা।"""
        ids = [to_int(i) for i in request.form.getlist("ids")]
        action = request.form.get("action")
        if not ids:
            flash("অন্তত একজন কর্মচারী টিক দিন।", "error")
        elif action == "activate":
            get_db().executemany("UPDATE employees SET status = 'active', separation_reason = NULL, "
                                 "separation_date = NULL WHERE id = ?", [(i,) for i in ids])
            get_db().commit()
            flash(f"{bn_num(len(ids))} জনকে আবার সক্রিয় করা হয়েছে।", "ok")
        else:
            reason = request.form.get("reason") or "অন্যান্য"
            when = clean_date(request.form.get("date")) or date.today().isoformat()
            get_db().executemany("UPDATE employees SET status = 'inactive', separation_reason = ?, separation_date = ? "
                                 "WHERE id = ?", [(reason, when, i) for i in ids])
            get_db().commit()
            flash(f"{bn_num(len(ids))} জনকে নিষ্ক্রিয় তালিকায় সরানো হয়েছে ({reason}, {bn_date(when)})। "
                  "পরের বেতন শিটে আসবেন না, লগইন বন্ধ; পুরনো রেকর্ড থাকবে।", "ok")
        return redirect(request.referrer or url_for("employees"))

    @app.route("/admin/departments", methods=["GET", "POST"])
    @admin_required
    def departments():
        """শাখা/বিভাগ: নতুন যোগ, নাম পরিবর্তন, খালি শাখা মোছা।"""
        conn = get_db()
        if request.method == "POST":
            action = request.form.get("action")
            name = re.sub(r"\s+", " ", request.form.get("name", "")).strip()
            old = request.form.get("old", "").strip()
            if action == "add" and name:
                if conn.execute("SELECT 1 FROM departments WHERE name = ?", (name,)).fetchone():
                    flash(f"'{name}' শাখা আগে থেকেই আছে।", "warn")
                else:
                    conn.execute("INSERT INTO departments (name) VALUES (?)", (name,))
                    flash(f"নতুন শাখা '{name}' যোগ হয়েছে। কর্মচারীর তথ্য সম্পাদনায় এখন এটি বাছাই করা যাবে।", "ok")
            elif action == "rename" and old and name and old != name:
                n = conn.execute("UPDATE employees SET department = ? WHERE department = ?", (name, old)).rowcount
                if conn.execute("SELECT 1 FROM departments WHERE name = ?", (name,)).fetchone():
                    conn.execute("DELETE FROM departments WHERE name = ?", (old,))  # দুই শাখা এক হয়ে গেল
                else:
                    conn.execute("UPDATE departments SET name = ? WHERE name = ?", (name, old))
                flash(f"'{old}' এর নাম '{name}' করা হয়েছে ({bn_num(n)} জন কর্মচারী)।", "ok")
            elif action == "delete" and old:
                n = conn.execute("SELECT COUNT(*) FROM employees WHERE department = ?", (old,)).fetchone()[0]
                if n:
                    flash(f"'{old}' শাখায় {bn_num(n)} জন কর্মচারী আছেন, তাই মোছা যাবে না। আগে তাঁদের অন্য শাখায় সরান।", "error")
                else:
                    conn.execute("DELETE FROM departments WHERE name = ?", (old,))
                    flash(f"'{old}' শাখা মুছে ফেলা হয়েছে।", "ok")
            conn.commit()
            return redirect(url_for("departments"))
        counts = dict(conn.execute("SELECT department, COUNT(*) FROM employees WHERE status = 'active' "
                                   "GROUP BY department").fetchall())
        totals = dict(conn.execute("SELECT department, COUNT(*) FROM employees GROUP BY department").fetchall())
        return render_template("admin/departments.html", depts=all_departments(), counts=counts, totals=totals)

    # ---------- বেতন স্কেল ও ভাতা

    @app.route("/admin/payscale", methods=["GET", "POST"])
    @admin_required
    def payscale_page():
        if request.method == "POST":
            for key in ["office_name", "office_address", "accounts_officer_name", "accounts_officer_designation",
                        "bank_rules"]:
                save_setting(key, request.form.get(key, "").strip())
            for key in ["festival_pct", "boishakhi_pct"]:
                save_setting(key, to_float(request.form.get(key)))
            get_db().commit()
            flash("সংরক্ষিত হয়েছে।", "ok")
            return redirect(url_for("payscale_page"))
        return render_template("admin/payscale.html", OLD=payscale.OLD, NEW=payscale.NEW, HR15=payscale.HR15,
                               HR26=payscale.HR26)

    # ---------- মাসিক বেতন শিট

    @app.route("/admin/payroll")
    @admin_required
    def payroll_list():
        runs = list(reversed(run_totals()))
        today = date.today()
        return render_template("admin/payroll_list.html", runs=runs, today=today, arrears_done=arrears_final(),
                               arrears_draft=arrears_already_given())

    @app.route("/admin/payroll/new", methods=["POST"])
    @admin_required
    def payroll_new():
        f = request.form
        year, month = to_int(f.get("year")), to_int(f.get("month"))
        kind = f.get("kind") or "salary"
        if not (2000 <= year <= 2100 and 1 <= month <= 12) or kind not in RUN_KINDS:
            flash("সঠিক মাস ও বছর দিন।", "error")
            return redirect(url_for("payroll_list"))
        conn = get_db()
        if kind == "salary":
            existing = conn.execute("SELECT id FROM payroll_runs WHERE year = ? AND month = ? AND kind = 'salary' "
                                    "AND manual = 0", (year, month)).fetchone()
            if existing:
                flash(f"{month_label(year, month)} এর বেতন শিট আগেই তৈরি হয়েছে।", "warn")
                return redirect(url_for("payroll_view", run_id=existing["id"]))
            manual = conn.execute("SELECT COUNT(*) FROM payslips p JOIN payroll_runs r ON r.id = p.run_id "
                                  "WHERE r.year = ? AND r.month = ? AND r.kind = 'salary' AND r.manual = 1",
                                  (year, month)).fetchone()[0]
            if manual:
                flash(f"{month_label(year, month)} এর বেতন {bn_num(manual)} জনের জন্য হাতে এন্ট্রি করা আছে "
                      "(পূর্ববর্তী মাসের বেতন এন্ট্রি)। একই মাসের বেতন দুবার যেন না আসে, তাই আগে সেগুলো মুছে তারপর শিট তৈরি করুন।",
                      "error")
                return redirect(url_for("payroll_list"))
        # বকেয়া একবারই: চূড়ান্ত হয়ে গেলে আর নয়; খসড়া থাকলে সেটাই খোলে
        if kind == "arrears" and (prev := arrears_already_given()):
            if prev["status"] == "final":
                flash(f"জুলাই–সেপ্টেম্বর ২০২৬ এর বকেয়া {month_label(prev['year'], prev['month'])} এর শিটে চূড়ান্তভাবে দেওয়া হয়েছে। "
                      "বকেয়া একবারই দেওয়া হয় — রেকর্ড শিটের তালিকায় আছে।", "error")
                return redirect(url_for("payroll_list"))
            flash("বকেয়ার খসড়া শিট আগেই তৈরি হয়েছে — সেটি যাচাই করে চূড়ান্ত করুন।", "warn")
            return redirect(url_for("payroll_view", run_id=prev["id"]))
        opts = {"kind": kind, "title": (f.get("title") or "").strip() or None}
        if kind == "festival":
            religion = f.get("religion") if f.get("religion") in RELIGIONS else None
            if not religion:
                flash("উৎসব ভাতার জন্য ধর্ম নির্বাচন করুন।", "error")
                return redirect(url_for("payroll_list"))
            opts.update(religion=religion,
                        multiplier=to_float(f.get("multiplier"), 0) or FESTIVAL_MULTIPLIER.get(religion, 1),
                        title=opts["title"] or f"উৎসব ভাতা ({religion})")
        elif kind == "other":
            pct, amount = to_float(f.get("pct"), 0), to_float(f.get("amount"), 0)
            if not opts["title"] or not (pct or amount):
                flash("অন্যান্য ভাতার নাম এবং মূল বেতনের % অথবা জনপ্রতি টাকা দিন।", "error")
                return redirect(url_for("payroll_list"))
            opts.update(pct=pct or None, amount=None if pct else amount)
        elif kind != "salary":
            opts["title"] = opts["title"] or RUN_KINDS[kind]
        run_id, skipped = generate_run(year, month, False, False, f.get("note", "").strip(), False, **opts)
        label = RUN_KINDS[kind] if kind == "salary" else opts["title"]
        n = conn.execute("SELECT COUNT(*) FROM payslips WHERE run_id = ?", (run_id,)).fetchone()[0]
        flash(f"{month_label(year, month)} — {label}: {bn_num(n)} জনের খসড়া শিট তৈরি হয়েছে। যাচাই করে 'চূড়ান্ত করুন' চাপুন।", "ok")
        if skipped:
            flash(f"{bn_num(len(skipped))} জন শিটে আসেননি: "
                  + ", ".join(f"{e['name_bn']} ({why})" for e, why in skipped[:20]), "warn")
        return redirect(url_for("payroll_view", run_id=run_id))

    def filtered_slips(run_id, args):
        """শিট দেখা ও Excel এর আগে ফিল্টার/সাজানো: শাখা ও গ্রেড (একাধিক), সাজানোর ক্রম, শাখাভিত্তিক ভাগ।"""
        slips = get_db().execute("SELECT * FROM payslips WHERE run_id = ? ORDER BY code", (run_id,)).fetchall()
        all_depts = list(dict.fromkeys(s["department"] or "শাখা উল্লেখ নেই" for s in slips))
        all_grades = sorted({s["grade"] for s in slips if s["grade"]})
        depts = [d for d in args.getlist("dept") if d]
        grades = [to_int(g) for g in args.getlist("grade") if g]
        if depts:
            slips = [s for s in slips if (s["department"] or "শাখা উল্লেখ নেই") in depts]
        if grades:
            slips = [s for s in slips if s["grade"] in grades]
        sort = args.get("sort", "code")
        keys = {"code": lambda s: s["code"] or "", "grade": lambda s: (s["grade"] or 99, s["code"] or ""),
                "name": lambda s: s["name_bn"] or "", "basic": lambda s: (-(s["basic"] or s["ref_basic"] or 0), s["code"]),
                "net": lambda s: (-s["net"], s["code"] or "")}
        slips = sorted(slips, key=keys.get(sort, keys["code"]))
        group = "1" in args.getlist("group") if "group" in args else True
        parts = []
        if depts:
            parts.append("শাখা: " + ", ".join(depts))
        if grades:
            parts.append("গ্রেড: " + ", ".join(bn_num(g) for g in sorted(grades)))
        return {"slips": slips, "all_depts": all_depts, "all_grades": all_grades, "sel_depts": depts,
                "sel_grades": grades, "sort": sort, "group": group, "filter_label": "; ".join(parts)}

    @app.route("/admin/payroll/<int:run_id>")
    @admin_required
    def payroll_view(run_id):
        run = get_run_or_404(run_id)
        fs = filtered_slips(run_id, request.args)
        slips = fs["slips"]
        totals = {f: round(sum(s[f] for s in slips), 2)
                  for f in ALLOWANCE_FIELDS + DEDUCTION_FIELDS + ["gross", "total_deduction", "net"]}
        on_prl = {e["id"] for e in get_db().execute("SELECT id, dob, prl_date FROM employees")
                  if (payscale.parse_date(e["prl_date"]) or payscale.prl_date(e["dob"]) or date.max)
                  <= date(run["year"], run["month"], 1)}
        missing = get_db().execute(
            "SELECT * FROM employees WHERE status = 'active' AND id NOT IN (SELECT employee_id FROM payslips WHERE run_id = ?)",
            (run_id,)).fetchall() if run["kind"] == "salary" else []
        prev = get_db().execute(
            "SELECT COALESCE(SUM(p.net),0) AS net FROM payroll_runs r JOIN payslips p ON p.run_id = r.id "
            "WHERE r.kind = 'salary' AND r.manual = 0 AND r.year * 12 + r.month = ? ", (run["year"] * 12 + run["month"] - 1,)).fetchone()
        grouped = group_by_department(slips) if fs["group"] else [(None, slips)]
        groups = [(d, g, {f: round(sum(s[f] for s in g), 2) for f in ["gross", "total_deduction", "net"]})
                  for d, g in grouped]
        banks = [(b, len(rows), round(sum(r["net"] for r in rows), 2)) for b, rows in bank_groups(run, slips)]
        return render_template("admin/payroll_view.html", run=run, slips=slips, totals=totals, missing=missing, banks=banks,
                               groups=groups, prev_net=prev["net"] if prev and run["kind"] == "salary" else 0,
                               on_prl=on_prl, fs=fs, query=request.query_string.decode())

    @app.route("/admin/payroll/<int:run_id>/excel")
    @admin_required
    def payroll_excel(run_id):
        run = get_run_or_404(run_id)
        fs = filtered_slips(run_id, request.args)
        name = "salary-sheet" if run["kind"] == "salary" else run["kind"]
        if run["kind"] == "arrears":
            return xlsx_response(excel.arrears_workbook(run, fs["slips"], get_settings(), group=fs["group"],
                                                        filter_label=fs["filter_label"]),
                                 f"arrears-{run['year']}-{run['month']:02d}.xlsx")
        # বেতন শিট: নাসিকের নিজস্ব ফরম্যাট (অঞ্চল/শাখা, কর্তন সমূহ) — ?fmt=simple দিলে সাধারণ তালিকা
        if run["kind"] == "salary" and request.args.get("fmt") != "simple":
            return xlsx_response(excel.ncc_workbook(run, fs["slips"], get_settings(), filter_label=fs["filter_label"]),
                                 f"{name}-{run['year']}-{run['month']:02d}.xlsx")
        return xlsx_response(excel.payroll_workbook(run, fs["slips"], get_settings(), group=fs["group"],
                                                    filter_label=fs["filter_label"]),
                             f"{name}-{run['year']}-{run['month']:02d}{'-list' if run['kind'] == 'salary' else ''}.xlsx")

    @app.route("/admin/payroll/<int:run_id>/bank-statement")
    @admin_required
    def bank_statement(run_id):
        """ব্যাংক স্টেটমেন্ট: bank=all → সব ব্যাংকের আলাদা Excel একসাথে (zip); bank=০,১… → একটি ব্যাংকের Excel।"""
        run = get_run_or_404(run_id)
        groups = bank_groups(run, filtered_slips(run_id, request.args)["slips"])
        settings = get_settings()
        ym = f"{run['year']}-{run['month']:02d}"
        label = f"{MONTHS_BN[run['month']]}-{bn_num(run['year'])}"
        what = "বেতন" if run["kind"] == "salary" else (run["title"] or "ভাতা")

        def fname(bank):
            return f"{what}-{bank or 'হিসাব নম্বর নেই'}-{label}.xlsx".replace(" ", "-").replace("/", "-")

        pick = request.args.get("bank", "all")
        if pick != "all":
            i = to_int(pick, -1)
            if not 0 <= i < len(groups):
                abort(404)
            bank, rows = groups[i]
            return xlsx_response(excel.bank_statement_workbook(run, rows, settings, bank), fname(bank))
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for bank, rows in groups:
                if rows:
                    z.writestr(fname(bank), excel.bank_statement_workbook(run, rows, settings, bank).getvalue())
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"bank-statement-{ym}.zip", mimetype="application/zip")

    @app.route("/admin/payroll/<int:run_id>/finalize", methods=["POST"])
    @admin_required
    def payroll_finalize(run_id):
        get_run_or_404(run_id)
        get_db().execute("UPDATE payroll_runs SET status = 'final', finalized_at = CURRENT_TIMESTAMP WHERE id = ?", (run_id,))
        get_db().commit()
        flash("বেতন শিট চূড়ান্ত হয়েছে। কর্মচারীরা এখন নিজ ড্যাশবোর্ডে বেতন বিবরণী দেখতে পাবেন।", "ok")
        return redirect(url_for("payroll_view", run_id=run_id))

    @app.route("/admin/payroll/<int:run_id>/reopen", methods=["POST"])
    @admin_required
    def payroll_reopen(run_id):
        run = get_run_or_404(run_id)
        if run["kind"] == "arrears" or run["include_arrears"]:
            flash("বকেয়ার শিট একবার চূড়ান্ত হলে আর খসড়ায় ফেরত নেওয়া যায় না — এটি স্থায়ী রেকর্ড।", "error")
            return redirect(url_for("payroll_view", run_id=run_id))
        get_db().execute("UPDATE payroll_runs SET status = 'draft', finalized_at = NULL WHERE id = ?", (run_id,))
        get_db().commit()
        flash("বেতন শিট আবার খসড়া অবস্থায় ফেরত গেছে।", "warn")
        return redirect(url_for("payroll_view", run_id=run_id))

    @app.route("/admin/payroll/<int:run_id>/delete", methods=["POST"])
    @admin_required
    def payroll_delete(run_id):
        run = get_run_or_404(run_id)
        if require_draft(run):
            get_db().execute("DELETE FROM payroll_runs WHERE id = ?", (run_id,))
            get_db().commit()
            flash(f"{month_label(run['year'], run['month'])} এর খসড়া বেতন শিট মুছে ফেলা হয়েছে।", "ok")
            return redirect(url_for("payroll_list"))
        return redirect(url_for("payroll_view", run_id=run_id))

    @app.route("/admin/payroll/<int:run_id>/recalc", methods=["POST"])
    @admin_required
    def payroll_recalc(run_id):
        """খসড়া শিটের সকলের বেতন বর্তমান তথ্য অনুযায়ী আবার হিসাব + নতুন যুক্ত কর্মচারী অন্তর্ভুক্ত।"""
        run = get_run_or_404(run_id)
        if require_draft(run):
            conn = get_db()
            for slip in conn.execute("SELECT * FROM payslips WHERE run_id = ?", (run_id,)).fetchall():
                regenerate_slip(run, slip)
            settings = get_settings()
            for emp in conn.execute("SELECT * FROM employees WHERE status = 'active' AND id NOT IN "
                                    "(SELECT employee_id FROM payslips WHERE run_id = ?)", (run_id,)).fetchall():
                s, _why = slip_for(emp, run, settings)
                if s:
                    insert_slip(run_id, emp, s)
            conn.commit()
            flash("সকলের বেতন বর্তমান তথ্য অনুযায়ী পুনরায় হিসাব করা হয়েছে।", "ok")
        return redirect(url_for("payroll_view", run_id=run_id))

    @app.route("/admin/payroll/<int:run_id>/slip/<int:slip_id>/remove", methods=["POST"])
    @admin_required
    def payslip_remove(run_id, slip_id):
        run = get_run_or_404(run_id)
        if require_draft(run):
            row = get_db().execute("SELECT employee_id FROM payslips WHERE id = ? AND run_id = ?", (slip_id, run_id)).fetchone()
            if row:
                record_loan_payments(run_id, row["employee_id"], [])
            get_db().execute("DELETE FROM payslips WHERE id = ? AND run_id = ?", (slip_id, run_id))
            get_db().commit()
            flash("এই মাসের শিট থেকে বাদ দেওয়া হয়েছে।", "ok")
        return redirect(url_for("payroll_view", run_id=run_id))

    @app.route("/admin/payroll/<int:run_id>/slip/<int:slip_id>/edit", methods=["GET", "POST"])
    @admin_required
    def payslip_edit(run_id, slip_id):
        run = get_run_or_404(run_id)
        slip = get_db().execute("SELECT * FROM payslips WHERE id = ? AND run_id = ?", (slip_id, run_id)).fetchone() or abort(404)
        if request.method == "POST":
            if not require_draft(run):
                return redirect(url_for("payroll_view", run_id=run_id))
            if request.form.get("action") == "recalc":
                if regenerate_slip(run, slip):
                    get_db().commit()
                    flash("কর্মচারীর বর্তমান তথ্য অনুযায়ী পুনরায় হিসাব করা হয়েছে।", "ok")
                else:
                    flash("কর্মচারীর গ্রেড/মূল বেতন অসম্পূর্ণ, তাই হিসাব করা যায়নি।", "error")
                return redirect(url_for("payroll_view", run_id=run_id))
            s = {f: round(to_float(request.form.get(f)), 2) for f in ALLOWANCE_FIELDS + DEDUCTION_FIELDS}
            for kind in ("other_allowance", "other_deduction"):
                names, amounts = request.form.getlist(f"{kind}_name"), request.form.getlist(f"{kind}_amount")
                loan_ids = request.form.getlist(f"{kind}_loan") or [""] * len(names)
                s[kind] = [{"name": n.strip(), "amount": round(to_float(a), 2), **({"loan_id": int(l)} if l else {})}
                           for n, a, l in zip(names, amounts, loan_ids) if n.strip()]
            finalize_totals(s)
            get_db().execute(
                f"UPDATE payslips SET {', '.join(f + ' = ?' for f in ALLOWANCE_FIELDS + DEDUCTION_FIELDS)}, "
                "other_allowance = ?, other_deduction = ?, gross = ?, total_deduction = ?, net = ? WHERE id = ?",
                [s[f] for f in ALLOWANCE_FIELDS + DEDUCTION_FIELDS] +
                [json.dumps(s["other_allowance"], ensure_ascii=False), json.dumps(s["other_deduction"], ensure_ascii=False),
                 s["gross"], s["total_deduction"], s["net"], slip_id])
            record_loan_payments(run_id, slip["employee_id"], s["other_deduction"])
            get_db().commit()
            flash(f"{slip['name_bn']} এর এ মাসের বেতন সংশোধন হয়েছে।", "ok")
            return redirect(url_for("payroll_view", run_id=run_id))
        return render_template("admin/payslip_edit.html", run=run, slip=slip, allowance_fields=ALLOWANCE_FIELDS,
                               deduction_fields=DEDUCTION_FIELDS)

    # ---------- রিপোর্ট

    @app.route("/admin/reports")
    @admin_required
    def reports():
        runs = run_totals()
        rows, prev = [], None
        for r in runs:
            is_salary = r["kind"] == "salary"
            rows.append({"run": r, "diff": round(r["net"] - prev["net"], 2) if prev and is_salary else None,
                         "hdiff": (r["headcount"] - prev["headcount"]) if prev and is_salary else None})
            if is_salary:
                prev = r
        years = {}
        for r in runs:
            y = years.setdefault(r["year"], {"gross": 0, "deduction": 0, "net": 0, "months": 0})
            y["gross"] += r["gross"]
            y["deduction"] += r["deduction"]
            y["net"] += r["net"]
            y["months"] += 1
        return render_template("admin/reports.html", rows=list(reversed(rows)), years=years,
                               chart=run_totals("salary")[-12:])

    @app.route("/admin/reports/excel")
    @admin_required
    def reports_excel():
        return xlsx_response(excel.history_workbook(run_totals(), get_settings()), "monthly-salary-summary.xlsx")

    # ---------- কর্মচারীর নিজস্ব ড্যাশবোর্ড

    def my_employee():
        user = current_user()
        if user["role"] != "employee" or not user["employee_id"]:
            abort(403)
        return get_db().execute("SELECT * FROM employees WHERE id = ?", (user["employee_id"],)).fetchone()

    @app.route("/me")
    @login_required
    def my_dashboard():
        emp = my_employee()
        slips = employee_slips(emp["id"], final_only=True)
        year = date.today().year
        ytd = round(sum(s["net"] for s in slips if s["year"] == year), 2)
        return render_template("me/dashboard.html", emp=emp, slips=slips, latest=slips[0] if slips else None,
                               loans=employee_loans(emp["id"]),
                               chart=list(reversed(slips[:12])), ytd=ytd, year=year, **pay_context(emp))

    @app.route("/me/profile")
    @login_required
    def my_profile():
        emp = my_employee()
        return render_template("me/profile.html", emp=emp, **pay_context(emp))

    @app.route("/me/profile/edit", methods=["GET", "POST"])
    @login_required
    def my_profile_edit():
        """কর্মচারী শুধু ব্যক্তিগত ও ব্যাংক তথ্য বদলাতে পারবেন; বেতন সংক্রান্ত কিছু নয়। প্রতিটি পরিবর্তনের রেকর্ড থাকে।"""
        emp = my_employee()
        if request.method == "POST":
            conn = get_db()
            changed = []
            for field, label in SELF_EDIT_FIELDS:
                new = (request.form.get(field) or "").strip() or None
                if field == "name_bn" and not new:
                    continue
                if new != emp[field]:
                    conn.execute(f"UPDATE employees SET {field} = ? WHERE id = ?", (new, emp["id"]))
                    conn.execute("INSERT INTO employee_changes (employee_id, field, old_value, new_value) "
                                 "VALUES (?, ?, ?, ?)", (emp["id"], label, emp[field], new))
                    changed.append(label)
            conn.commit()
            flash(("হালনাগাদ হয়েছে: " + ", ".join(changed)) if changed else "কোনো পরিবর্তন হয়নি।", "ok")
            return redirect(url_for("my_profile"))
        return render_template("me/profile_edit.html", emp=emp, fields=SELF_EDIT_FIELDS)

    @app.route("/me/history.xlsx")
    @login_required
    def my_history_xlsx():
        emp = my_employee()
        slips = list(reversed(employee_slips(emp["id"], final_only=True)))
        return xlsx_response(excel.employee_history_workbook(emp, slips, get_settings()), f"salary-{emp['code']}.xlsx")

    @app.route("/me/statement")
    @login_required
    def my_statement():
        return statement(my_employee()["id"])

    @app.route("/statement/<int:emp_id>")
    @login_required
    def statement(emp_id):
        """ব্যক্তির বেতন বিবরণী: নির্দিষ্ট মাস থেকে মাস পর্যন্ত, শেষে হিসাবরক্ষণ কর্মকর্তার সইয়ের ঘর; PDF ডাউনলোড।"""
        user = current_user()
        if user["role"] != "admin" and user["employee_id"] != emp_id:
            abort(403)
        emp = get_db().execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone() or abort(404)
        all_slips = list(reversed(employee_slips(emp_id, final_only=user["role"] != "admin")))
        months = [(s["year"], s["month"]) for s in all_slips]

        def ym(arg, default):
            m = re.fullmatch(r"(\d{4})-(\d{1,2})", request.args.get(arg, ""))
            return (int(m.group(1)), int(m.group(2))) if m else default
        start = ym("from", months[0] if months else (date.today().year, 1))
        end = ym("to", months[-1] if months else (date.today().year, date.today().month))
        if start > end:
            start, end = end, start
        kinds = [k for k in request.args.getlist("kind") if k in RUN_KINDS] or list(RUN_KINDS)
        detail = request.args.get("detail") == "1"
        slips = [s for s in all_slips if start <= (s["year"], s["month"]) <= end and s["kind"] in kinds]
        rows = []
        for s in slips:
            oa, od = json.loads(s["other_allowance"]), json.loads(s["other_deduction"])
            rows.append({"s": s, "other_a": round(sum(float(x["amount"]) for x in oa), 2),
                         "loan": round(sum(float(x["amount"]) for x in od if x.get("loan_id")), 2),
                         "other_d": round(sum(float(x["amount"]) for x in od if not x.get("loan_id")), 2)})
        totals = {f: round(sum(s[f] for s in slips), 2)
                  for f in ALLOWANCE_FIELDS + DEDUCTION_FIELDS + ["gross", "total_deduction", "net"]}
        for k in ("other_a", "loan", "other_d"):
            totals[k] = round(sum(r[k] for r in rows), 2)
        # বিস্তারিত: যে ভাতাগুলো এই সময়ে কারো আছে শুধু সেগুলোর কলাম
        acols = [(f, SHORT_LABELS[f]) for f in ALLOWANCE_FIELDS if totals[f]] if detail else []
        years = sorted({y for y, _ in months} | {date.today().year, 2026})
        employees = get_db().execute("SELECT id, code, name_bn, department FROM employees ORDER BY code").fetchall() \
            if user["role"] == "admin" else []
        return render_template("statement.html", emp=emp, slips=slips, rows=rows, totals=totals, start=start, end=end,
                               years=years, today=date.today(), kinds=kinds, detail=detail, acols=acols,
                               employees=employees)

    @app.route("/admin/statements")
    @admin_required
    def statements():
        """যেকোনো কর্মচারীর বেতন বিবরণী: শাখা → কর্মচারী, মাসের সীমা, শিটের ধরন, বিস্তারিত/সংক্ষিপ্ত।"""
        emp_id = to_int(request.args.get("emp"), None)
        if emp_id:
            args = request.args.to_dict(flat=False)
            args.pop("emp", None)
            return redirect(url_for("statement", emp_id=emp_id, **args))
        employees = get_db().execute("SELECT id, code, name_bn, designation, department, status FROM employees "
                                     "ORDER BY code").fetchall()
        today = date.today()
        return render_template("admin/statements.html", employees=employees, depts=all_departments(),
                               years=sorted({2026, today.year, today.year - 1}), today=today)

    @app.route("/payslip/<int:slip_id>")
    @login_required
    def payslip(slip_id):
        slip = get_db().execute(
            "SELECT p.*, r.year, r.month, r.status, r.kind, r.title FROM payslips p JOIN payroll_runs r ON r.id = p.run_id WHERE p.id = ?",
            (slip_id,)).fetchone() or abort(404)
        user = current_user()
        if user["role"] != "admin" and (slip["employee_id"] != user["employee_id"] or slip["status"] != "final"):
            abort(403)
        emp = get_db().execute("SELECT * FROM employees WHERE id = ?", (slip["employee_id"],)).fetchone()
        return render_template("payslip.html", slip=slip, emp=emp)

    @app.errorhandler(400)
    def bad_request(_e):
        msg = ("অনুরোধটি যাচাই করা যায়নি। পাতাটি রিফ্রেশ করে আবার চেষ্টা করুন। বারবার হলে ঠিকানাটি "
               "https:// দিয়ে খুলুন — নিরাপত্তার কারণে http:// এ লগইন কাজ করে না।")
        return render_template("error.html", code=400, message=msg), 400

    @app.errorhandler(403)
    def forbidden(_e):
        return render_template("error.html", code=403, message="এই পাতাটি দেখার অনুমতি আপনার নেই।"), 403

    @app.errorhandler(404)
    def not_found(_e):
        return render_template("error.html", code=404, message="পাতাটি পাওয়া যায়নি।"), 404


app = create_app()

if __name__ == "__main__":
    app.run(host=os.environ.get("HOST", "127.0.0.1"), port=int(os.environ.get("PORT", 5000)),
            debug=os.environ.get("FLASK_DEBUG") == "1")
