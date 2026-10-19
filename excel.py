"""Excel (.xlsx) ফাইল তৈরি — মাসিক বেতন শিট, ব্যাংক অ্যাডভাইস, মাসভিত্তিক সারাংশ, কর্মচারী ইমপোর্ট টেমপ্লেট।"""
import io
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from bn import MONTHS_BN, bn_num, bn_taka
from payroll import group_by_department

# নিকস (Nikosh) — বাংলাদেশের সরকারি অফিসের কম্পিউটারে প্রায় সবসময় থাকে; না থাকলে Excel নিজে অন্য বাংলা ফন্ট বসায়
FONT = "Nikosh"
HEAD_FILL = PatternFill("solid", fgColor="1F5F4A")
TOTAL_FILL = PatternFill("solid", fgColor="E6F0EC")
DEPT_FILL = PatternFill("solid", fgColor="F3F7F5")
GRAND_FILL = PatternFill("solid", fgColor="CFE3DA")
THIN = Side(style="thin", color="B7C4BE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
# বাংলা অঙ্কে (১২,৩৪,৫৬৭) দেখানোর ফরম্যাট — মান সংখ্যাই থাকে, তাই যোগফল/সূত্র ঠিকমতো কাজ করে
BN = "[$-5000445]"


def _lakh(dec=""):
    return (f"{BN}[>=10000000]##\\,##\\,##\\,##0{dec};{BN}[>=100000]##\\,##\\,##0{dec};{BN}##,##0{dec}")


MONEY = _lakh()
MONEY_PAISA = _lakh(".00")
BN_INT = f"{BN}0"


LOGO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static", "logo.png")


def _logo(ws, size=56, anchor="A1"):
    """শিটের উপরে বাম কোণে নাসিকের লোগো (Pillow না থাকলে লোগো ছাড়াই)।"""
    try:
        from openpyxl.drawing.image import Image
        img = Image(LOGO)
    except Exception:
        return
    img.width = img.height = size
    ws.add_image(img, anchor)


def _title(ws, office, title, ncols):
    _logo(ws)
    ws.row_dimensions[1].height = 22
    for row, (text, size) in enumerate([(office["office_name"], 15), (office["office_address"], 11), (title, 13)], 1):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        c = ws.cell(row=row, column=1, value=text)
        c.font = Font(name=FONT, size=size, bold=row != 2)
        c.alignment = Alignment(horizontal="center")


def _header(ws, row, headers):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 36


def _style_body(ws, first, last, ncols, money_from):
    for r in range(first, last + 1):
        for col in range(1, ncols + 1):
            c = ws.cell(row=r, column=col)
            c.font = Font(name=FONT, bold=c.font.bold)
            c.border = BORDER
            v = c.value
            if col >= money_from:
                c.number_format = MONEY_PAISA if isinstance(v, float) and not v.is_integer() else MONEY
            elif isinstance(v, int):
                c.number_format = BN_INT


def _totals_row(ws, row, first, last, label_col, money_cols, label="সর্বমোট", ncols=None):
    ws.cell(row=row, column=label_col, value=label).font = Font(name=FONT, bold=True)
    for col in range(1, (ncols or ws.max_column) + 1):
        ws.cell(row=row, column=col).fill = TOTAL_FILL
        ws.cell(row=row, column=col).border = BORDER
    for col in money_cols:
        L = get_column_letter(col)
        c = ws.cell(row=row, column=col, value=f"=SUM({L}{first}:{L}{last})")
        c.font = Font(name=FONT, bold=True)
        paisa = any(isinstance(ws.cell(row=r, column=col).value, float)
                    and not ws.cell(row=r, column=col).value.is_integer() for r in range(first, last + 1))
        c.number_format = MONEY_PAISA if paisa else MONEY


def _signature(ws, row, ncols, settings):
    """নিচে সইয়ের ঘর: বামে প্রস্তুতকারী, ডানে হিসাবরক্ষণ কর্মকর্তা (নাম ও পদবি সেটিং থেকে)।"""
    right = max(ncols - 3, 4)
    lines_left = ["……………………………", "প্রস্তুতকারী"]
    lines_right = ["……………………………", settings.get("accounts_officer_name") or "",
                   settings.get("accounts_officer_designation") or "হিসাবরক্ষণ কর্মকর্তা", settings["office_name"]]
    for i, t in enumerate(lines_left):
        ws.cell(row=row + i, column=2, value=t).font = Font(name=FONT, bold=i == 1)
    for i, t in enumerate(x for x in lines_right if x):
        ws.cell(row=row + i, column=right, value=t).font = Font(name=FONT, bold=i == 1)


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _save(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def payroll_workbook(run, slips, settings, group=True, filter_label=""):
    wb = Workbook()
    ws = wb.active
    salary = run["kind"] == "salary"
    ws.title = "বেতন শিট" if salary else "ভাতা শিট"
    period = f"{MONTHS_BN[run['month']]} {bn_num(run['year'])}"
    sheet_title = (f"মাসিক বেতন শিট — {period}" if salary else f"{run['title']} — {period}") + \
        (f" ({filter_label})" if filter_label else "")

    # যে ভাতা এ মাসে কারো নেই সেই কলাম বাদ
    all_cols = [("basic", "মূল বেতন"), ("house_rent", "বাড়িভাড়া"), ("medical", "চিকিৎসা"),
                ("education", "শিক্ষা সহায়ক"), ("tiffin", "টিফিন"), ("transport", "যাতায়াত"), ("mobile", "মোবাইল"),
                ("wash", "ধোলাই"), ("special", "বিশেষ সুবিধা (মহার্ঘ)"), ("pf_office", "অফিস প্রদত্ত PF"),
                ("festival", "উৎসব ভাতা"), ("boishakhi", "নববর্ষ ভাতা")]
    forced = ("basic", "house_rent", "medical") if salary else ()
    allowance_cols = [(k, h) for k, h in all_cols if k in forced or any(s[k] for s in slips)]
    has_pf = any(s["pf_fund"] or s["pf_own"] for s in slips)
    has_stamp = any(s["revenue_stamp"] for s in slips)
    ded_cols = ([("pf_fund", "PF (অফিস অংশ জমা)"), ("pf_own", "PF (নিজ অংশ)")] if has_pf else []) + \
        ([("revenue_stamp", "রাজস্ব স্ট্যাম্প")] if has_stamp else [])

    info = ["ক্রম", "আইডি", "নাম", "পদবি", "গ্রেড", "হিসাব নং"] + ([] if salary else ["ভিত্তি মূল বেতন"])
    headers = (info + [h for _, h in allowance_cols] + ["অন্যান্য ভাতা / বকেয়া", "মোট বেতন"] + [h for _, h in ded_cols]
               + ["অন্যান্য কর্তন", "মোট কর্তন", "নিট প্রদেয়"])
    n = len(headers)
    money_from = len(info) + 1
    _title(ws, settings, sheet_title, n)
    hr = 5
    _header(ws, hr, headers)

    # শাখাভিত্তিক: শাখার নাম → কর্মচারীরা → উপমোট → একটি ফাঁকা সারি
    r = hr
    subtotal_rows = []
    for dept, members in (group_by_department(slips) if group else [(None, slips)]):
        if dept is not None:
            r += 1
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
            c = ws.cell(row=r, column=1, value=f"{dept}  ({bn_num(len(members))} জন)")
            c.font = Font(name=FONT, bold=True, size=12, color="1F5F4A")
            c.fill = DEPT_FILL
            c.border = BORDER
        first = r + 1
        for i, s in enumerate(members, 1):
            r += 1
            oa = round(sum(float(x["amount"]) for x in json.loads(s["other_allowance"])), 2)
            od = round(sum(float(x["amount"]) for x in json.loads(s["other_deduction"])), 2)
            values = ([i, s["code"], s["name_bn"], s["designation"], s["grade"], s["account_no"]]
                      + ([] if salary else [s["ref_basic"]]) + [s[k] for k, _ in allowance_cols] + [oa, s["gross"]] + [s[k] for k, _ in ded_cols]
                      + [od, s["total_deduction"], s["net"]])
            for col, v in enumerate(values, 1):
                ws.cell(row=r, column=col, value=v)
        _style_body(ws, first, r, n, money_from)
        if not salary:
            for rr in range(first, r + 1):
                ws.cell(row=rr, column=len(info)).number_format = MONEY
        r += 1
        _totals_row(ws, r, first, r - 1, 3, range(money_from, n + 1), label="উপমোট" if dept is not None else "মোট", ncols=n)
        subtotal_rows.append(r)
        r += 1  # শাখার মাঝে বিরতি
    r += 1
    ws.cell(row=r, column=3, value=f"সর্বমোট ({bn_num(len(slips))} জন)").font = Font(name=FONT, bold=True, size=12)
    for col in range(1, n + 1):
        cell = ws.cell(row=r, column=col)
        cell.fill = GRAND_FILL
        cell.border = BORDER
        if col >= money_from and subtotal_rows:
            L = get_column_letter(col)
            cell.value = "=" + "+".join(f"{L}{x}" for x in subtotal_rows)
            cell.font = Font(name=FONT, bold=True, size=12)
            cell.number_format = MONEY_PAISA
    _signature(ws, r + 4, n, settings)
    ws.freeze_panes = f"D{hr + 1}"
    _widths(ws, [6, 11, 26, 22, 7, 18] + ([] if salary else [14]) + [12] * (n - len(info)))
    ws.print_title_rows = f"{hr}:{hr}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ব্যাংক অ্যাডভাইস
    bank = wb.create_sheet("ব্যাংক অ্যাডভাইস")
    bh = ["ক্রম", "আইডি", "নাম", "পদবি", "হিসাব নং", "নিট প্রদেয় (টাকা)"]
    _title(bank, settings, f"ব্যাংক অ্যাডভাইস — {period}" + ("" if salary else f" ({run['title']})"), len(bh))
    _header(bank, hr, bh)
    for i, s in enumerate(slips, 1):
        for col, v in enumerate([i, s["code"], s["name_bn"], s["designation"], s["account_no"], s["net"]], 1):
            bank.cell(row=hr + i, column=col, value=v)
    blast = hr + max(len(slips), 1)
    _style_body(bank, hr + 1, blast, len(bh), 6)
    _totals_row(bank, blast + 1, hr + 1, blast, 3, [6])
    _signature(bank, blast + 5, len(bh), settings)
    _widths(bank, [6, 11, 26, 22, 22, 18])

    # শাখাভিত্তিক সারাংশ
    summary = wb.create_sheet("শাখাভিত্তিক সারাংশ")
    sh = ["শাখা", "কর্মচারী সংখ্যা", "মোট বেতন", "মোট কর্তন", "নিট প্রদেয়"]
    _title(summary, settings, f"শাখাভিত্তিক সারাংশ — {period}", len(sh))
    _header(summary, hr, sh)
    by_dept = {}
    for dept, group in group_by_department(slips):
        by_dept[dept] = [len(group), round(sum(x["gross"] for x in group), 2),
                         round(sum(x["total_deduction"] for x in group), 2), round(sum(x["net"] for x in group), 2)]
    for i, (dept, vals) in enumerate(by_dept.items(), 1):
        for col, v in enumerate([dept] + vals, 1):
            summary.cell(row=hr + i, column=col, value=v)
    slast = hr + max(len(by_dept), 1)
    _style_body(summary, hr + 1, slast, len(sh), 3)
    _totals_row(summary, slast + 1, hr + 1, slast, 1, range(2, 6))
    _signature(summary, slast + 5, len(sh), settings)
    _widths(summary, [40, 16, 18, 16, 18])
    return _save(wb)


def history_workbook(runs, settings):
    """মাসভিত্তিক বেতন প্রদানের সারাংশ ও পূর্ববর্তী মাসের সাথে পার্থক্য।"""
    wb = Workbook()
    ws = wb.active
    ws.title = "মাসভিত্তিক সারাংশ"
    headers = ["মাস", "অবস্থা", "কর্মচারী সংখ্যা", "মোট বেতন", "মোট কর্তন", "নিট প্রদান", "আগের মাসের তুলনায় পার্থক্য"]
    _title(ws, settings, "মাসভিত্তিক বেতন প্রদানের হিসাব", len(headers))
    _header(ws, 5, headers)
    prev = None
    for i, run in enumerate(runs, 1):
        diff = run["net"] - prev if prev is not None else None
        vals = [f"{MONTHS_BN[run['month']]} {bn_num(run['year'])}", "চূড়ান্ত" if run["status"] == "final" else "খসড়া",
                run["headcount"], run["gross"], run["deduction"], run["net"], diff]
        for col, v in enumerate(vals, 1):
            ws.cell(row=5 + i, column=col, value=v)
        prev = run["net"]
    last = 5 + max(len(runs), 1)
    _style_body(ws, 6, last, len(headers), 4)
    _widths(ws, [18, 10, 16, 18, 16, 18, 24])
    return _save(wb)


def employee_history_workbook(emp, slips, settings):
    wb = Workbook()
    ws = wb.active
    ws.title = "বেতন ইতিহাস"
    headers = ["মাস", "মূল বেতন", "বাড়িভাড়া", "চিকিৎসা", "শিক্ষা", "টিফিন", "যাতায়াত", "মোবাইল/ধোলাই/মহার্ঘ/PF",
               "উৎসব/নববর্ষ", "অন্যান্য ভাতা/বকেয়া", "মোট বেতন", "মোট কর্তন", "নিট প্রাপ্তি"]
    _title(ws, settings, f"বেতন ইতিহাস — {emp['name_bn']} ({emp['code']})", len(headers))
    _header(ws, 5, headers)
    for i, s in enumerate(slips, 1):
        oa = round(sum(float(x["amount"]) for x in json.loads(s["other_allowance"])), 2)
        others = s["mobile"] + s["wash"] + s["special"] + s["pf_office"]
        vals = [f"{MONTHS_BN[s['month']]} {bn_num(s['year'])}", s["basic"], s["house_rent"], s["medical"], s["education"],
                s["tiffin"], s["transport"], others, s["festival"] + s["boishakhi"], oa, s["gross"], s["total_deduction"],
                s["net"]]
        for col, v in enumerate(vals, 1):
            ws.cell(row=5 + i, column=col, value=v)
    last = 5 + max(len(slips), 1)
    _style_body(ws, 6, last, len(headers), 2)
    _totals_row(ws, last + 1, 6, last, 1, range(2, len(headers) + 1))
    _widths(ws, [18] + [13] * (len(headers) - 1))
    return _save(wb)


# কর্মচারী ইমপোর্ট/এক্সপোর্ট কলাম: (ডাটাবেস ফিল্ড, শিরোনাম)
EMPLOYEE_COLUMNS = [
    ("code", "আইডি*"), ("name_bn", "নাম (বাংলা)*"), ("name_en", "নাম (ইংরেজি)"), ("designation", "পদবি"),
    ("department", "শাখা/বিভাগ"), ("grade", "গ্রেড"), ("basic", "মূল বেতন (৩০-০৬-২০২৬)"),
    ("inc_2026", "ইনক্রিমেন্ট ০১-০৭-২০২৬ (হ্যাঁ/না)"), ("inc_2027", "ইনক্রিমেন্ট ০১-০৭-২০২৭ (হ্যাঁ/না)"),
    ("new_basic_override", "নতুন মূল বেতন ০১-০৭-২০২৬ (ঐচ্ছিক)"), ("pf", "PF প্রযোজ্য (হ্যাঁ/না)"),
    ("children", "সন্তান সংখ্যা (০/১/২)"), ("tiffin", "টিফিন ভাতা (হ্যাঁ/না)"), ("transport", "যাতায়াত ভাতা (হ্যাঁ/না)"),
    ("wash", "ধোলাই ভাতা (হ্যাঁ/না)"), ("govt_quarter", "সরকারি বাসা (হ্যাঁ/না)"),
    ("special_rate", "মহার্ঘ হার (auto/0.10/0.15/0)"),
    ("join_date", "যোগদানের তারিখ (YYYY-MM-DD)"), ("dob", "জন্ম তারিখ (YYYY-MM-DD)"),
    ("prl_date", "পিআরএল তারিখ (YYYY-MM-DD)"), ("employment_type", "চাকরির ধরন"), ("status", "অবস্থা (active/inactive)"),
    ("religion", "ধর্ম (ইসলাম/হিন্দু/বৌদ্ধ/খ্রিস্টান/অন্যান্য)"),
    ("father_name", "পিতার নাম"), ("mother_name", "মাতার নাম"), ("gender", "লিঙ্গ"), ("nid", "এনআইডি"),
    ("mobile", "মোবাইল"), ("email", "ইমেইল"), ("address", "ঠিকানা"), ("tin", "টিআইএন"),
    ("bank_name", "ব্যাংক"), ("bank_branch", "শাখা (ব্যাংক)"), ("account_no", "হিসাব নং"),
    ("ovr_mobile", "মোবাইল ভাতা (নির্দিষ্ট, ঐচ্ছিক)"),
    ("allow_items", "অতিরিক্ত ভাতা (নাম=টাকা; নাম=টাকা)"), ("ded_items", "নির্দিষ্ট কর্তন (নাম=টাকা; নাম=টাকা)"),
] + [(f"loan{i}_{part}", f"{kind} — {label}") for i, kind in enumerate(["গৃহ নির্মাণ ঋণ", "ভবিষ্য তহবিল অগ্রিম", "অন্যান্য ঋণ"])
     for part, label in [("inst", "কিস্তি (টাকা)"), ("total", "মোট কিস্তি"), ("paid", "পরিশোধিত কিস্তি")]]
IMPORT_LOAN_KINDS = ["গৃহ নির্মাণ ঋণ", "ভবিষ্য তহবিল অগ্রিম", "অন্যান্য ঋণ"]
FLAG_COLUMNS = {"inc_2026", "inc_2027", "pf", "tiffin", "transport", "wash", "govt_quarter"}


def employees_workbook(employees, settings, template=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "কর্মচারী"
    _header(ws, 1, [h for _, h in EMPLOYEE_COLUMNS])
    if template:
        sample = {"code": "NCC-0001", "name_bn": "মোঃ রহিম উদ্দিন", "name_en": "Md. Rahim Uddin", "gender": "পুরুষ",
                  "dob": "1980-05-12", "mobile": "01700000000", "designation": "অফিস সহকারী",
                  "department": "সাধারণ শাখা", "employment_type": "স্থায়ী", "join_date": "2008-07-01", "grade": 14,
                  "basic": 22380, "inc_2026": 1, "inc_2027": 1, "pf": 0, "children": 2, "tiffin": 1, "transport": 1,
                  "wash": 0, "govt_quarter": 0, "special_rate": "auto", "status": "active",
                  "bank_name": "সোনালী ব্যাংক", "bank_branch": "নারায়ণগঞ্জ কর্পোরেট", "account_no": "0000000000000"}
        employees = [sample]
    for r, emp in enumerate(employees, 2):
        for col, (field, _) in enumerate(EMPLOYEE_COLUMNS, 1):
            v = emp[field] if field in emp.keys() else None
            if field in FLAG_COLUMNS and v is not None:
                v = "হ্যাঁ" if v else "না"
            ws.cell(row=r, column=col, value=v).font = Font(name=FONT)
    for i in range(1, len(EMPLOYEE_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.freeze_panes = "C2"
    return _save(wb)


# ---------------------------------------------------------------------------
# নাসিকের নিজস্ব বেতন শিট ফরম্যাট (জুলাই'২০২৬ এর শিটের মতো): সারসংক্ষেপ নোট, সারাংশ, বিস্তারিত
# ---------------------------------------------------------------------------
REGIONS = ["নারায়ণগঞ্জ অঞ্চল", "সিদ্ধিরগঞ্জ অঞ্চল", "কদমরসূল অঞ্চল"]
LETTERS = "কখগঘঙচছজ"
# বিস্তারিত শিটের কলাম (A থেকে W)
NCC_COLS = ["নং", "নাম", "পদবী", "মূল বেতন", "বাড়ী ভাড়া", "চিকিৎসা", "টিফিন", "যাতায়াত", "শিক্ষা", "বকেয়া/ধোলাই",
            "বিশেষ সুবিধা", "ভবিষ্য তহঃ", "মোট বেতন", "গৃহ ঋণ", "অন্যান্য ঋণ", "বাসা ভাড়া কর্তন", "বিভিন্ন কর্তন", "আয়কর",
            "অফিস", "নিজ", "ভবিষ্য অগ্রিম", "মোট কর্তন", "নীট বেতন"]
# শীর্ষ কর্মকর্তা (সচিব/প্রধান নির্বাহী) — আলাদা কলাম (A থেকে P)
TOP_COLS = ["নং", "নাম", "পদবী", "মূল বেতন", "বাড়ী ভাড়া", "চিকিৎসা", "শিক্ষা", "মোবাইল", "ইন্টারনেট ও টেলিফোন",
            "গাড়ী রক্ষণাবেক্ষণ", "বিশেষ সুবিধা", "অন্যান্য ভাতা", "মোট", "আয়কর কর্তন", "অন্যান্য কর্তন", "নীট প্রাপ্য"]
# শীর্ষ কর্মকর্তার কলাম → বিস্তারিত শিটের কলাম (সর্বমোটে যোগ করার জন্য)
TOP_TO_NCC = {"D": "D", "E": "E", "F": "F", "G": "I", "H": "J", "I": "J", "J": "J", "K": "K", "L": "J", "M": "M",
              "N": "R", "O": "Q", "P": "W"}
TOP_DESIGNATIONS = ("সচিব", "প্রধান নির্বাহী কর্মকর্তা")


def region_of(dept):
    d = dept or ""
    if "সিদ্ধিরগঞ্জ" in d:
        return REGIONS[1]
    if "কদমরসূল" in d or "কদমরসুল" in d:
        return REGIONS[2]
    return REGIONS[0]


def is_top_officer(s):
    return (s["designation"] or "").strip() in TOP_DESIGNATIONS


def _items_of(js):
    return [x for x in json.loads(js or "[]") if float(x.get("amount") or 0)]


def _has(name, *words):
    return any(w in (name or "") for w in words)


def ncc_breakdown(s):
    """একটি বেতন বিবরণীকে নাসিক শিটের কলামে ভাগ করা। কলামগুলোর যোগফল সবসময় মোট বেতন/মোট কর্তনের সমান।"""
    r2 = lambda x: round(float(x or 0), 2)  # noqa: E731
    d = {"basic": r2(s["basic"]), "house_rent": r2(s["house_rent"]), "medical": r2(s["medical"]),
         "tiffin": r2(s["tiffin"]), "transport": r2(s["transport"]), "education": r2(s["education"]),
         "special": r2(s["special"]), "pf_office": r2(s["pf_office"]), "mobile": r2(s["mobile"]),
         "internet": 0.0, "car": 0.0}
    for x in _items_of(s["other_allowance"]):
        if _has(x["name"], "ইন্টারনেট", "টেলিফোন"):
            d["internet"] += r2(x["amount"])
        elif _has(x["name"], "গাড়ি", "গাড়ী"):
            d["car"] += r2(x["amount"])
    listed = sum(d[k] for k in ("basic", "house_rent", "medical", "tiffin", "transport", "education", "special",
                                "pf_office"))
    # ধোলাই, বকেয়া, মোবাইল ও অন্যান্য সব ভাতা একসাথে "বকেয়া/ধোলাই" কলামে
    d["misc_allow"] = round(r2(s["gross"]) - listed, 2)
    d["top_other"] = round(d["misc_allow"] - d["mobile"] - d["internet"] - d["car"], 2)
    ded = {"house_loan": 0.0, "other_loan": 0.0, "rent": 0.0, "misc": 0.0, "tax": 0.0,
           "pf_fund": r2(s["pf_fund"]), "pf_own": r2(s["pf_own"]), "pf_adv": 0.0}
    for x in _items_of(s["other_deduction"]):
        n, a = x["name"], r2(x["amount"])
        if _has(n, "গৃহ নির্মাণ", "গৃহ ঋণ", "গৃহঋণ"):
            ded["house_loan"] += a
        elif _has(n, "ভবিষ্য তহবিল অগ্রিম", "ভবিষ্য অগ্রিম", "পিএফ অগ্রিম", "PF অগ্রিম"):
            ded["pf_adv"] += a
        elif _has(n, "ঋণ"):
            ded["other_loan"] += a
        elif _has(n, "বাসা ভাড়া", "বাড়ি ভাড়া কর্তন", "বাড়ী ভাড়া কর্তন"):
            ded["rent"] += a
        elif _has(n, "আয়কর"):
            ded["tax"] += a
        else:
            ded["misc"] += a
    ded["misc"] += r2(s["revenue_stamp"])
    d.update(ded)
    d["gross"], d["total_deduction"], d["net"] = r2(s["gross"]), r2(s["total_deduction"]), r2(s["net"])
    return d


def _v(x):
    """০ হলে ঘর খালি (মূল শিটের মতো)।"""
    if not x:
        return None
    return int(x) if float(x).is_integer() else round(x, 2)


def _cell(ws, row, col, value=None, bold=False, size=10, fill=None, align=None, fmt=None, border=True, wrap=False):
    c = ws.cell(row=row, column=col)
    if value is not None:
        c.value = value
    c.font = Font(name=FONT, bold=bold, size=size)
    if fill:
        c.fill = fill
    if border:
        c.border = BORDER
    if align or wrap:
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fmt:
        c.number_format = fmt
    return c


def _money_fmt(ws, row, first_col, last_col):
    for col in range(first_col, last_col + 1):
        ws.cell(row=row, column=col).number_format = MONEY_PAISA if _paisa(ws.cell(row=row, column=col).value) else MONEY


def _paisa(v):
    return isinstance(v, float) and not v.is_integer()


def _ncc_header(ws, r, cols, ded_from=None, pf_cols=None):
    """দুই সারির শিরোনাম; ভবিষ্য তহবিল কর্তনের নিচে অফিস/নিজ।"""
    n = len(cols)
    for col, h in enumerate(cols, 1):
        if pf_cols and col in pf_cols:
            continue
        ws.merge_cells(start_row=r, start_column=col, end_row=r + 1, end_column=col)
        _cell(ws, r, col, h, bold=True, fill=HEAD_FILL, align="center", wrap=True).font = \
            Font(name=FONT, bold=True, color="FFFFFF", size=10)
        _cell(ws, r + 1, col, fill=HEAD_FILL)
    if pf_cols:
        a, b = pf_cols
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        _cell(ws, r, a, "ভবিষ্য তহবিল কর্তন", fill=HEAD_FILL, align="center", wrap=True).font = \
            Font(name=FONT, bold=True, color="FFFFFF", size=10)
        _cell(ws, r, b, fill=HEAD_FILL)
        for col in (a, b):
            _cell(ws, r + 1, col, cols[col - 1], fill=HEAD_FILL, align="center").font = \
                Font(name=FONT, bold=True, color="FFFFFF", size=10)
    ws.row_dimensions[r].height = 30
    return n


def ncc_workbook(run, slips, settings, filter_label=""):
    """নাসিকের মাসিক বেতন শিট — অঞ্চল ও শাখাভিত্তিক বিস্তারিত, সারাংশ ও সারসংক্ষেপ নোট (সব বাংলায়)।"""
    period = f"{MONTHS_BN[run['month']]}'{bn_num(run['year'])} খ্রিঃ"
    top = [s for s in slips if is_top_officer(s)]
    rest = [s for s in slips if not is_top_officer(s)]
    regions = {name: [] for name in REGIONS}
    bs = [ncc_breakdown(s) for s in slips]
    # কোন কলামে পয়সা আছে — শুধু সেখানেই যোগফলে দুই ঘর দশমিক
    det_keys = dict(zip("DEFGHIJKLMNOPQRSTUVW", [
        ["basic"], ["house_rent"], ["medical"], ["tiffin"], ["transport"], ["education"], ["misc_allow"], ["special"],
        ["pf_office"], ["gross"], ["house_loan"], ["other_loan"], ["rent"], ["misc"], ["tax"], ["pf_fund"], ["pf_own"],
        ["pf_adv"], ["total_deduction"], ["net"]]))

    def fmt(*letters):
        keys = [k for L in letters for k in det_keys.get(L, [])]
        return MONEY_PAISA if any(_paisa(float(b[k])) for b in bs for k in keys) else MONEY

    for dept, members in group_by_department(rest):
        regions[region_of(dept)].append((dept, members))

    wb = Workbook()
    note = wb.active
    note.title = "সারসংক্ষেপ নোট"
    summary = wb.create_sheet("সারাংশ")
    ws = wb.create_sheet("বিস্তারিত")
    DET = "'বিস্তারিত'!"

    # ---------------- বিস্তারিত ----------------
    n = len(NCC_COLS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    _cell(ws, 1, 1, settings["office_name"], bold=True, size=14, align="center", border=False)
    _logo(ws, size=50)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 20
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    _cell(ws, 2, 1, f"{period} মাসের বেতন-ভাতার বিবরণী" + (f" ({filter_label})" if filter_label else ""),
          bold=True, size=12, align="center", border=False)
    r = 3
    letter = iter(LETTERS)
    top_rows = []          # শীর্ষ কর্মকর্তাদের সারি (TOP_COLS বিন্যাসে)
    dept_totals = []       # (অঞ্চল, শাখা, মোট সারি)
    region_totals = {}     # অঞ্চল → মোট সারি
    serial = 0

    if top:
        r += 1
        _cell(ws, r, 1, f"{next(letter)})", bold=True, border=False)
        _cell(ws, r, 2, "শীর্ষ কর্মকর্তা", bold=True, border=False)
        r += 1
        _ncc_header(ws, r, TOP_COLS)
        r += 1
        for s in top:
            r += 1
            serial += 1
            b = ncc_breakdown(s)
            vals = [serial, s["name_bn"], s["designation"], b["basic"], b["house_rent"], b["medical"], b["education"],
                    b["mobile"], b["internet"], b["car"], b["special"], b["top_other"]]
            for col, v in enumerate(vals, 1):
                _cell(ws, r, col, _v(v) if col > 3 else v, wrap=col in (2, 3))
            _cell(ws, r, 13, f"=SUM(D{r}:L{r})", bold=True)
            _cell(ws, r, 14, _v(b["tax"]))
            _cell(ws, r, 15, _v(round(b["total_deduction"] - b["tax"], 2)))
            _cell(ws, r, 16, f"=M{r}-N{r}-O{r}", bold=True)
            ws.cell(row=r, column=1).number_format = BN_INT
            _money_fmt(ws, r, 4, 16)
            top_rows.append(r)
        r += 1

    for region in REGIONS:
        depts = regions[region]
        if not depts:
            continue
        r += 1
        _cell(ws, r, 1, f"{next(letter)})", bold=True, size=12, border=False)
        _cell(ws, r, 2, region, bold=True, size=12, border=False)
        rows_in_region = []
        for di, (dept, members) in enumerate(depts, 1):
            r += 1
            if dept != region:
                _cell(ws, r, 1, di, bold=True, border=False, fmt=BN_INT)
                _cell(ws, r, 2, f"{dept}  ({bn_num(len(members))} জন)", bold=True, border=False)
            ws.merge_cells(start_row=r, start_column=14, end_row=r, end_column=n)
            _cell(ws, r, 14, "কর্তন সমূহ", bold=True, align="center")
            r += 1
            _ncc_header(ws, r, NCC_COLS, pf_cols=(19, 20))
            r += 1
            first = r + 1
            for s in members:
                r += 1
                serial += 1
                b = ncc_breakdown(s)
                vals = [serial, s["name_bn"], s["designation"]] + [_v(b[k]) for k in (
                    "basic", "house_rent", "medical", "tiffin", "transport", "education", "misc_allow", "special",
                    "pf_office")]
                for col, v in enumerate(vals, 1):
                    _cell(ws, r, col, v, wrap=col in (2, 3))
                _cell(ws, r, 13, f"=SUM(D{r}:L{r})", bold=True)
                for col, k in enumerate(("house_loan", "other_loan", "rent", "misc", "tax", "pf_fund", "pf_own",
                                         "pf_adv"), 14):
                    _cell(ws, r, col, _v(b[k]))
                _cell(ws, r, 22, f"=SUM(N{r}:U{r})")
                _cell(ws, r, 23, f"=M{r}-V{r}", bold=True)
                ws.cell(row=r, column=1).number_format = BN_INT
                _money_fmt(ws, r, 4, n)
            r += 1
            _cell(ws, r, 2, "মোট", bold=True, fill=TOTAL_FILL)
            for col in range(1, n + 1):
                c = _cell(ws, r, col, fill=TOTAL_FILL, bold=True)
                if col >= 4:
                    L = get_column_letter(col)
                    c.value = f"=SUM({L}{first}:{L}{r - 1})"
                    c.number_format = MONEY_PAISA if any(_paisa(ws.cell(row=x, column=col).value)
                                                         for x in range(first, r)) else MONEY
            dept_totals.append((region, dept, r))
            rows_in_region.append(r)
            r += 1  # শাখার মাঝে বিরতি
        if len(rows_in_region) > 1:
            r += 1
            _cell(ws, r, 2, f"{region}ের মোট", bold=True, fill=DEPT_FILL)
            for col in range(1, n + 1):
                c = _cell(ws, r, col, fill=DEPT_FILL, bold=True)
                if col >= 4:
                    L = get_column_letter(col)
                    c.value = "=" + "+".join(f"{L}{x}" for x in rows_in_region)
                    c.number_format = fmt(L)
            r += 1
        region_totals[region] = r - 1 if len(rows_in_region) > 1 else rows_in_region[0]

    # সর্বমোট
    r += 1
    labels = "+".join(LETTERS[i] for i in range(len(region_totals) + (1 if top else 0)))
    _cell(ws, r, 2, f"সর্বমোট ({labels}) — {bn_num(len(slips))} জন", bold=True, size=11, fill=GRAND_FILL)
    grand_row = r
    for col in range(1, n + 1):
        c = _cell(ws, r, col, fill=GRAND_FILL, bold=True, size=11)
        if col >= 4:
            L = get_column_letter(col)
            parts = [f"{L}{x}" for x in region_totals.values()]
            parts += [f"{tc}{tr}" for tr in top_rows for tc, nc in TOP_TO_NCC.items() if nc == L]
            if L == "V":  # শীর্ষ কর্মকর্তার মোট কর্তন
                parts += [f"N{tr}+O{tr}" for tr in top_rows]
            c.value = "=" + "+".join(parts) if parts else 0
            c.number_format = fmt(L)
    _signature(ws, r + 4, n, settings)
    _widths(ws, [5, 22, 18] + [10.5] * 9 + [11.5] + [9.5] * 9 + [11, 12])
    ws.freeze_panes = "D3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True

    # ---------------- সারাংশ (শাখাভিত্তিক মোট) ----------------
    sc = ["নং", "বিবরণ"] + NCC_COLS[3:]          # A..V; বিস্তারিত D..W → সারাংশ C..V
    sn = len(sc)
    summary.merge_cells(start_row=1, start_column=1, end_row=1, end_column=sn)
    _cell(summary, 1, 1, settings["office_name"], bold=True, size=14, align="center", border=False)
    _logo(summary, size=50)
    summary.row_dimensions[1].height = 24
    summary.row_dimensions[2].height = 20
    summary.merge_cells(start_row=2, start_column=1, end_row=2, end_column=sn)
    _cell(summary, 2, 1, f"{period} মাসের বেতন-ভাতার সারাংশ", bold=True, size=12, align="center", border=False)
    r = 4
    summary.merge_cells(start_row=r, start_column=13, end_row=r, end_column=sn)
    _cell(summary, r, 13, "কর্তন সমূহ", bold=True, align="center")
    r += 1
    _ncc_header(summary, r, sc, pf_cols=(18, 19))
    r += 1
    det_to_sum = lambda L: get_column_letter(ord(L) - ord("A"))  # noqa: E731  বিস্তারিত D → সারাংশ C
    sum_region_rows = []
    letter = iter(LETTERS)
    if top:
        r += 1
        _cell(summary, r, 1, f"{next(letter)})", bold=True)
        _cell(summary, r, 2, ", ".join(s["designation"] for s in top), bold=True, wrap=True)
        for col in range(3, sn + 1):
            L = get_column_letter(col + 1)  # বিস্তারিত শিটের কলাম
            parts = [f"{DET}{tc}{tr}" for tr in top_rows for tc, nc in TOP_TO_NCC.items() if nc == L]
            if L == "V":
                parts += [f"{DET}N{tr}+{DET}O{tr}" for tr in top_rows]
            _cell(summary, r, col, "=" + "+".join(parts) if parts else None, fmt=fmt(L))
        sum_region_rows.append(r)
    for region in REGIONS:
        rows = [(d, tr) for reg, d, tr in dept_totals if reg == region]
        if not rows:
            continue
        r += 1
        _cell(summary, r, 1, f"{next(letter)})", bold=True, size=11, border=False)
        _cell(summary, r, 2, region, bold=True, size=11, border=False)
        first = r + 1
        for i, (dept, tr) in enumerate(rows, 1):
            r += 1
            _cell(summary, r, 1, i, fmt=BN_INT)
            _cell(summary, r, 2, region if dept == region else dept, wrap=True)
            for col in range(3, sn + 1):
                _cell(summary, r, col, f"={DET}{get_column_letter(col + 1)}{tr}", fmt=fmt(get_column_letter(col + 1)))
        r += 1
        _cell(summary, r, 2, "মোট", bold=True, fill=TOTAL_FILL)
        for col in range(1, sn + 1):
            c = _cell(summary, r, col, fill=TOTAL_FILL, bold=True)
            if col >= 3:
                L = get_column_letter(col)
                c.value = f"=SUM({L}{first}:{L}{r - 1})"
                c.number_format = fmt(get_column_letter(col + 1))
        sum_region_rows.append(r)
        r += 1
    r += 1
    _cell(summary, r, 2, f"সর্বমোট ({labels})", bold=True, size=11, fill=GRAND_FILL)
    for col in range(1, sn + 1):
        c = _cell(summary, r, col, fill=GRAND_FILL, bold=True, size=11)
        if col >= 3:
            c.value = f"={DET}{get_column_letter(col + 1)}{grand_row}"
            c.number_format = fmt(get_column_letter(col + 1))
    _signature(summary, r + 4, sn, settings)
    _widths(summary, [5, 34] + [11] * (sn - 2))
    summary.freeze_panes = "C7"
    summary.page_setup.orientation = "landscape"
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 0
    summary.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------------- সারসংক্ষেপ নোট (অঞ্চলভিত্তিক) ----------------
    nh = ["বিবরণ", "মোট বেতন ভাতাদি", "ভবিষ্য তহবিল কর্তন", "বিভিন্ন ঋণের কিস্তি ও অন্যান্য কর্তন", "বাসা ভাড়া কর্তন",
          "আয়কর কর্তন", "মোট কর্তন", "নীট বেতন", "আনুতোষিক"]
    note.merge_cells("A1:I1")
    _cell(note, 1, 1, settings["office_name"], bold=True, size=14, align="center", border=False)
    _logo(note, size=50)
    note.row_dimensions[1].height = 30
    note.merge_cells("A2:I2")
    _cell(note, 2, 1, f"{settings['office_name']}ে কর্মরত কর্মকর্তা ও কর্মচারীদের {period} মাসের বেতন-ভাতা পরিশোধের "
                      "নিমিত্তে বিবরণ নিম্নে উপস্থাপন করা হলো -", size=11, wrap=True, border=False)
    note.row_dimensions[2].height = 36
    for col, h in enumerate(nh, 1):
        _cell(note, 4, col, h, align="center", wrap=True).font = Font(name=FONT, bold=True, color="FFFFFF")
        note.cell(row=4, column=col).fill = HEAD_FILL
    for col, h in enumerate(["", "১", "২", "৩", "৪", "৫", "৬ (২-৫)", "৭ (১-৬)", ""], 1):
        _cell(note, 5, col, h or None, align="center", size=9)
    note.row_dimensions[4].height = 45
    # সারাংশ শিটের কলাম: C মূল … L মোট বেতন, M গৃহ ঋণ, N অন্যান্য ঋণ, O বাসা ভাড়া, P বিভিন্ন, Q আয়কর, R/S PF, T অগ্রিম, U মোট কর্তন, V নীট
    S = "'সারাংশ'!"
    note_fmt = {2: fmt("M"), 3: fmt("S", "T", "U"), 4: fmt("N", "O", "Q"), 5: fmt("P"), 6: fmt("R"), 7: fmt("V"),
                8: fmt("W")}
    r = 5
    region_names = ([", ".join(s["designation"] for s in top)] if top else []) + \
        [reg for reg in REGIONS if any(x[0] == reg for x in dept_totals)]
    for name, sr in zip(region_names, sum_region_rows):
        r += 1
        _cell(note, r, 1, name, wrap=True)
        _cell(note, r, 2, f"={S}L{sr}", fmt=note_fmt[2])
        _cell(note, r, 3, f"={S}R{sr}+{S}S{sr}+{S}T{sr}", fmt=note_fmt[3])
        _cell(note, r, 4, f"={S}M{sr}+{S}N{sr}+{S}P{sr}", fmt=note_fmt[4])
        _cell(note, r, 5, f"={S}O{sr}", fmt=note_fmt[5])
        _cell(note, r, 6, f"={S}Q{sr}", fmt=note_fmt[6])
        _cell(note, r, 7, f"=SUM(C{r}:F{r})", fmt=note_fmt[7])
        _cell(note, r, 8, f"=B{r}-G{r}", fmt=note_fmt[8])
        _cell(note, r, 9, fmt=MONEY_PAISA)
    last = r
    r += 1
    _cell(note, r, 1, "সর্বমোট", bold=True, fill=GRAND_FILL)
    for col in range(2, 10):
        L = get_column_letter(col)
        _cell(note, r, col, f"=SUM({L}6:{L}{last})", bold=True, fill=GRAND_FILL, fmt=note_fmt.get(col, MONEY_PAISA))
    total_row = r

    # সংক্ষিপ্ত হিসাব (টাকার অঙ্ক পাইথনেই হিসাব করে লেখা, যাতে সূত্র ছাড়াও পড়া যায়)
    tot = lambda *ks: round(sum(b[k] for b in bs for k in ks), 2)  # noqa: E731
    lines = [("বেতন ভাতা", tot("net")), ("ভবিষ্য তহবিল", tot("pf_fund", "pf_own", "pf_adv")),
             ("বিভিন্ন ঋণের কিস্তি ও অন্যান্য কর্তন বাবদ", tot("house_loan", "other_loan", "misc")),
             ("বাসা ভাড়া কর্তন", tot("rent")), ("আয়কর", tot("tax"))]
    r = total_row + 2
    for label, amount in lines:
        note.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        _cell(note, r, 1, f"{label} = {bn_taka(amount)}/- টাকা", size=11, border=False)
        r += 1
    note.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    _cell(note, r, 1, f"মোট = {bn_taka(tot('gross'))}/- টাকা", bold=True, size=11, border=False)
    _signature(note, r + 4, 9, settings)
    _widths(note, [34, 15, 14, 16, 12, 12, 14, 15, 13])
    note.page_setup.orientation = "landscape"
    note.sheet_properties.pageSetUpPr.fitToPage = True
    note.page_setup.fitToWidth = 1
    note.page_setup.fitToHeight = 0
    return _save(wb)
