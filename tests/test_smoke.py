"""সম্পূর্ণ কাজের প্রবাহের স্মোক টেস্ট:  python tests/test_smoke.py"""
import io
import os
import re
import sqlite3
import subprocess
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from openpyxl import load_workbook  # noqa: E402

from app import create_app  # noqa: E402
from excel import EMPLOYEE_COLUMNS  # noqa: E402

COLS = [f for f, _ in EMPLOYEE_COLUMNS]


def csrf(client):
    with client.session_transaction() as s:
        return s.get("csrf")


def post(client, url, data=None, **kw):
    client.get("/login")  # নিশ্চিত করে session এ csrf আছে
    d = dict(data or {})
    d["csrf"] = csrf(client)
    return client.post(url, data=d, **kw)


def test_full_flow():
    db = os.path.join(tempfile.mkdtemp(), "t.db")
    app = create_app(db)
    c = app.test_client()
    from db import get_db

    assert c.get("/admin").status_code == 302
    r = post(c, "/login", {"username": "admin", "password": "admin123"})
    assert c.get("/admin").headers["Location"].endswith("/password")  # প্রথম লগইনে পাসওয়ার্ড বদল বাধ্যতামূলক
    r = post(c, "/password", {"old": "admin123", "new": "Admin#2026", "confirm": "Admin#2026"}, follow_redirects=True)
    assert "ড্যাশবোর্ড" in r.get_data(as_text=True)

    # নতুন কর্মচারী: ক্যালকুলেটরের নমুনা (গ্রেড ১৩, ১৯,৮১০), PF প্রযোজ্য, ২ সন্তান
    r = post(c, "/admin/employees/new", {"code": "E1", "name_bn": "পরীক্ষা কর্মচারী", "grade": "13", "basic": "১৯,৮১০",
                                         "children": "2", "dob": "1970-01-01", "status": "active", "department": "সাধারণ শাখা",
                                         "inc_2026": "1", "inc_2027": "1", "tiffin": "1", "transport": "1", "pf": "1",
                                         "special_rate": "auto"}, follow_redirects=True)
    html = r.get_data(as_text=True)
    assert "২৮,৫০৫" in html and "৩২,০০২.৫০" in html and "৩৭,৩০০" in html  # তিন কিস্তি
    assert "বকেয়া: জুলাই–সেপ্টেম্বর ২০২৬" in html
    emp_id = int(re.search(r"/admin/employees/(\d+)/edit", html).group(1))
    post(c, f"/admin/employees/{emp_id}/items", {"kind": "deduction", "name": "ঋণ কিস্তি", "amount": "১০০০"})

    # Excel ইমপোর্ট (টেমপ্লেটে নমুনা + আরেকজন, ইনক্রিমেন্ট 'না')
    wb = load_workbook(io.BytesIO(c.get("/admin/employees/template.xlsx").data))
    ws = wb.active
    for field, v in {"code": "E2", "name_bn": "দ্বিতীয় কর্মচারী", "grade": 14, "basic": 10200, "inc_2026": "না",
                     "department": "হিসাব বিভাগ"}.items():
        ws.cell(row=3, column=COLS.index(field) + 1, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = post(c, "/admin/employees/import", {"file": (buf, "e.xlsx")}, content_type="multipart/form-data",
             follow_redirects=True)
    assert "নতুন ২ জন" in r.get_data(as_text=True)
    with app.app_context():
        e2 = get_db().execute("SELECT * FROM employees WHERE code = 'E2'").fetchone()
        assert (e2["inc_2026"], e2["inc_2027"], e2["tiffin"], e2["pf"]) == (0, 1, 1, 0)

    # অক্টোবর ২০২৬ শিট: উৎসব ভাতা + বকেয়া
    r = post(c, "/admin/payroll/new", {"year": "2026", "month": "10", "festival": "1", "arrears": "1"},
             follow_redirects=True)
    html = r.get_data(as_text=True)
    assert "বেতন শিট — অক্টোবর ২০২৬" in html
    run_id = int(re.search(r"/admin/payroll/(\d+)/excel", html).group(1))
    slip_id = int(re.search(rf"/admin/payroll/{run_id}/slip/(\d+)/edit", html).group(1))

    import payscale
    with app.app_context():
        emp = get_db().execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone()
        s = get_db().execute("SELECT * FROM payslips WHERE employee_id = ?", (emp_id,)).fetchone()
        a = payscale.arrears(emp)
        assert (s["basic"], s["house_rent"], s["medical"], s["education"], s["tiffin"], s["transport"],
                s["festival"], s["pf_office"], s["special"]) == (28505, 8914.5, 1500, 1000, 200, 300, 28505, 2850.5, 0)
        assert s["phase"] == "১ম কিস্তি"
        gross = 28505 + 8914.5 + 1500 + 1000 + 200 + 300 + 28505 + 2850.5 + a["total"] + a["pf_total"]
        ded = 2850.5 * 2 + 1000 + 2 * a["pf_total"]
        assert round(s["gross"], 2) == round(gross, 2) and round(s["net"], 2) == round(gross - ded, 2)

    r = c.get(f"/admin/payroll/{run_id}/excel")
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["বেতন শিট", "ব্যাংক অ্যাডভাইস", "শাখাভিত্তিক সারাংশ"]
    assert "PF (নিজ অংশ)" in [cell.value for cell in wb["বেতন শিট"][5]]

    # একজনের বেতন হাতে সংশোধন (পয়সাসহ)
    form = {f: "0" for f in ["basic", "house_rent", "medical", "education", "tiffin", "transport", "mobile", "wash",
                             "special", "pf_office", "festival", "boishakhi", "pf_fund", "pf_own", "revenue_stamp"]}
    form.update({"basic": "20000.50", "other_allowance_name": "বকেয়া", "other_allowance_amount": "500"})
    post(c, f"/admin/payroll/{run_id}/slip/{slip_id}/edit", form)
    with app.app_context():
        s = get_db().execute("SELECT * FROM payslips WHERE id = ?", (slip_id,)).fetchone()
        assert s["gross"] == 20500.5 and s["net"] == 20500.5

    # কর্মচারী: খসড়া দেখতে পারবে না, নিজের তথ্য সীমিতভাবে বদলাতে পারবে
    emp_c = app.test_client()
    post(emp_c, "/login", {"username": "E1", "password": "E1"})
    assert emp_c.get("/me").headers["Location"].endswith("/password")
    post(emp_c, "/password", {"old": "E1", "new": "secret99", "confirm": "secret99"})
    html = emp_c.get("/me").get_data(as_text=True)
    assert "এখনো কোনো বেতনের রেকর্ড নেই" in html and "মূল বেতন নির্ধারণ: ধাপে ধাপে" in html
    assert emp_c.get("/admin").status_code == 403
    assert emp_c.get(f"/admin/employees/{emp_id}/edit").status_code == 403
    assert emp_c.get("/me/profile/edit").status_code == 200
    post(emp_c, "/me/profile/edit", {"name_bn": "পরীক্ষা কর্মচারী", "account_no": "123456789", "basic": "99999",
                                     "grade": "1", "mobile": "01711111111"})
    with app.app_context():
        e = get_db().execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone()
        assert (e["account_no"], e["mobile"], e["basic"], e["grade"]) == ("123456789", "01711111111", 19810, 13)
        assert get_db().execute("SELECT COUNT(*) FROM employee_changes WHERE employee_id = ?", (emp_id,)).fetchone()[0] == 2
    assert "123456789" in c.get("/admin").get_data(as_text=True)  # অ্যাডমিন ড্যাশবোর্ডে পরিবর্তন দেখা যায়

    post(c, f"/admin/payroll/{run_id}/finalize")
    assert "অক্টোবর ২০২৬" in emp_c.get("/me").get_data(as_text=True)
    with app.app_context():
        own = get_db().execute("SELECT id FROM payslips WHERE employee_id = ?", (emp_id,)).fetchone()["id"]
        other = get_db().execute("SELECT id FROM payslips WHERE employee_id != ?", (emp_id,)).fetchone()["id"]
    assert emp_c.get(f"/payslip/{own}").status_code == 200
    assert emp_c.get(f"/payslip/{other}").status_code == 403
    assert emp_c.get("/me/history.xlsx").status_code == 200
    assert emp_c.get("/me/profile").status_code == 200

    for url in ["/admin", "/admin/employees", f"/admin/employees/{emp_id}", f"/admin/employees/{emp_id}/edit",
                "/admin/employees/new", "/admin/payroll", "/admin/reports", "/admin/reports/excel", "/admin/payscale",
                "/admin/employees/export.xlsx", f"/admin/employees/{emp_id}/history.xlsx", "/admin/employees/import",
                f"/admin/payroll/{run_id}", f"/payslip/{own}"]:
        assert c.get(url).status_code == 200, url

    # শাখার নাম একসাথে বদল
    post(c, "/admin/departments/rename", {"old": "সাধারণ শাখা", "new": "প্রশাসন"})
    with app.app_context():
        assert get_db().execute("SELECT department FROM employees WHERE id = ?", (emp_id,)).fetchone()[0] == "প্রশাসন"

    # চূড়ান্ত শিট মোছা যাবে না; CSRF ছাড়া POST বাতিল
    post(c, f"/admin/payroll/{run_id}/delete")
    with app.app_context():
        assert get_db().execute("SELECT 1 FROM payroll_runs WHERE id = ?", (run_id,)).fetchone()
    assert c.post(f"/admin/payroll/{run_id}/delete").status_code == 400

    # জানুয়ারি ২০২৮: নতুন হারে ভাতা
    post(c, "/admin/payroll/new", {"year": "2028", "month": "1"})
    with app.app_context():
        s = get_db().execute("SELECT p.* FROM payslips p JOIN payroll_runs r ON r.id = p.run_id "
                             "WHERE r.year = 2028 AND p.employee_id = ?", (emp_id,)).fetchone()
        assert (s["basic"], s["house_rent"], s["medical"], s["tiffin"], s["transport"], s["mobile"]) == \
            (37300, 14920, 4000, 500, 600, 150)

    # ৫ বার ভুল পাসওয়ার্ডের পর লগইন বন্ধ
    bad = app.test_client()
    for _ in range(5):
        assert post(bad, "/login", {"username": "admin", "password": "wrong"}).status_code == 200
    assert post(bad, "/login", {"username": "admin", "password": "Admin#2026"}).status_code == 429


def test_old_database_is_upgraded():
    """আগের সংস্করণের ডাটাবেসে নতুন কলাম নিজে থেকে যোগ হয়।"""
    old_schema = subprocess.run(["git", "show", "1783f0e:schema.sql"], cwd=ROOT, capture_output=True, text=True)
    if old_schema.returncode != 0:
        return
    db = os.path.join(tempfile.mkdtemp(), "old.db")
    conn = sqlite3.connect(db)
    conn.executescript(old_schema.stdout)
    conn.executemany("INSERT INTO pay_grades VALUES (?, 1, 1)", [(g,) for g in range(1, 21)])  # পুরনো সংস্করণ এগুলো রাখত
    conn.execute("INSERT INTO employees (code, name_bn, grade, basic) VALUES ('OLD1', 'পুরনো', 16, 10260)")
    conn.commit()
    conn.close()
    app = create_app(db)
    with app.app_context():
        from db import get_db
        cols = {r[1] for r in get_db().execute("PRAGMA table_info(employees)")}
        assert {"inc_2026", "pf", "special_rate", "wash"} <= cols
        e = get_db().execute("SELECT * FROM employees WHERE code = 'OLD1'").fetchone()
        assert e["inc_2026"] == 1 and e["tiffin"] == 1


if __name__ == "__main__":
    test_full_flow()
    test_old_database_is_upgraded()
    print("সব টেস্ট সফল")
